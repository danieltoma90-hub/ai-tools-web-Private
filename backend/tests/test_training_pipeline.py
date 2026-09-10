# -*- coding: utf-8 -*-
"""Ce se întâmplă când specificația nu poate fi citită.

Cazul e real: o cheie API revocată face apelul să pice cu 401. Înainte, eroarea
era înghițită și utilizatorul primea o agendă „reușită" fără nicio particularitate
a clientului — exact ce trebuia să aducă tool-ul.
"""
from pathlib import Path
from unittest.mock import patch

import pytest

from ai_errors import mesaj_eroare_ai
from pipelines.training_pipeline import run_training_pipeline


class _Anthropic401(Exception):
    status_code = 401

    def __str__(self):
        return "Error code: 401 - {'type': 'authentication_error', 'message': 'API key is invalid.'}"


def _spec(tmp_path: Path) -> Path:
    from docx import Document

    p = tmp_path / "spec.docx"
    d = Document()
    d.add_paragraph("Rețeta se preia din laborator.")
    d.save(str(p))
    return p


def _module_spec(n: int = 3) -> list[dict]:
    return [
        {
            "nr": i + 1,
            "nume": f"Modul {i + 1} din specificație",
            "ore": 2.0,
            "audienta": "Operatori",
            "esential": False,
            "grupe": [("Conținut", [f"Subiect {i + 1}"])],
        }
        for i in range(n)
    ]


async def test_productie_esueaza_cand_specificatia_nu_poate_fi_citita(tmp_path):
    with patch(
        "pipelines.training_pipeline.construieste_module_din_spec",
        side_effect=_Anthropic401(),
    ):
        with pytest.raises(RuntimeError) as e:
            await run_training_pipeline(
                tip="productie", zile=3, client="ACME",
                spec_path=_spec(tmp_path), api_key="sk-x",
            )
    mesaj = str(e.value)
    assert "specificația" in mesaj
    assert "invalidă" in mesaj or "revocată" in mesaj


async def test_productie_fara_specificatie_nu_are_din_ce_construi():
    with pytest.raises(RuntimeError) as e:
        await run_training_pipeline(tip="productie", zile=3, client="ACME")
    assert "specificația clientului" in str(e.value)


async def test_productie_foloseste_doar_modulele_din_specificatie(tmp_path):
    """Nimic din catalogul CORE nu are ce căuta într-o agendă de producție."""
    with patch(
        "pipelines.training_pipeline.construieste_module_din_spec",
        return_value=_module_spec(3),
    ):
        docx, xlsx, sumar = await run_training_pipeline(
            tip="productie", zile=2, client="Solaris",
            spec_path=_spec(tmp_path), api_key="sk-x",
        )
    try:
        from docx import Document

        text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
        assert "Modul 1 din specificație" in text
        assert "Modul Depozit" not in text
        assert "Modul Contabilitate" not in text
        assert "specificația clientului" in text  # nota de sursă din antet
        assert sumar["module"] == 3
        assert sumar["ore_referinta"] == 6.0
        assert sumar["particularitati"] == 0
    finally:
        docx.unlink(missing_ok=True)
        xlsx.unlink(missing_ok=True)


async def test_productie_cu_specificatie_saraca_spune_de_ce(tmp_path):
    with patch(
        "pipelines.training_pipeline.construieste_module_din_spec",
        side_effect=ValueError("Din specificație nu a rezultat niciun modul de training."),
    ):
        with pytest.raises(ValueError) as e:
            await run_training_pipeline(
                tip="productie", zile=3, client="ACME",
                spec_path=_spec(tmp_path), api_key="sk-x",
            )
    assert "niciun modul" in str(e.value)


async def test_core_livreaza_standardul_dar_avertizeaza(tmp_path):
    with patch(
        "pipelines.training_pipeline.extrage_particularitati",
        side_effect=_Anthropic401(),
    ):
        docx, xlsx, sumar = await run_training_pipeline(
            tip="core", zile=3, client="ACME",
            spec_path=_spec(tmp_path), api_key="sk-x",
        )
    try:
        assert sumar["particularitati"] == 0
        assert sumar["avertisment"]
        assert "nu au putut fi extrase" in sumar["avertisment"]
        assert sumar["total_ore"] == 17.0  # programul standard e intact
    finally:
        docx.unlink(missing_ok=True)
        xlsx.unlink(missing_ok=True)


async def test_fara_specificatie_nu_exista_avertisment():
    docx, xlsx, sumar = await run_training_pipeline(tip="core", zile=3, client="ACME")
    try:
        assert sumar["avertisment"] == ""
    finally:
        docx.unlink(missing_ok=True)
        xlsx.unlink(missing_ok=True)


async def test_particularitatile_gasite_ajung_in_sumar(tmp_path):
    gasite = [
        {"modul": "Modul Depozit", "titlu": "Notă de cântar", "detaliu": "", "cod": ""},
        {"modul": "Modul Depozit", "titlu": "Retichetare", "detaliu": "", "cod": ""},
    ]
    with patch(
        "pipelines.training_pipeline.extrage_particularitati", return_value=gasite
    ):
        docx, xlsx, sumar = await run_training_pipeline(
            tip="core", zile=3, client="ACME",
            spec_path=_spec(tmp_path), api_key="sk-x",
        )
    try:
        assert sumar["particularitati"] == 2
        assert sumar["avertisment"] == ""
    finally:
        docx.unlink(missing_ok=True)
        xlsx.unlink(missing_ok=True)


@pytest.mark.parametrize(
    "exc,asteptat",
    [
        (_Anthropic401(), "invalidă"),
        (Exception("Your credit balance is too low"), "credite"),
        (type("E", (Exception,), {"status_code": 429})(), "limita"),
    ],
)
def test_mesajele_de_eroare_spun_ce_e_de_facut(exc, asteptat):
    assert asteptat in mesaj_eroare_ai(exc)

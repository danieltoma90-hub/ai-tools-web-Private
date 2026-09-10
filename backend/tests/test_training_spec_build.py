# -*- coding: utf-8 -*-
"""Construirea programului de Producție din specificație.

Regula pe care o apără aceste teste: în agendă intră doar ce scrie în document.
Un modul inventat îl pune pe consultant să predea ceva ce clientul nu are.
"""
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from skills.training.spec_build import (
    MAX_MODULE,
    _curata_modul,
    construieste_module_din_spec,
)


def _spec(tmp_path, text="Fluxul de producție trece prin Moară, Presă și Ambalare."):
    from docx import Document

    p = tmp_path / "spec.docx"
    d = Document()
    d.add_paragraph(text)
    d.save(str(p))
    return p


def _raspuns(payload: str):
    resp = MagicMock()
    resp.content = [MagicMock(text=payload)]
    client = MagicMock()
    client.messages.create = AsyncMock(return_value=resp)
    return client


PAYLOAD_BUN = """{"module": [
  {"nume": "Configurări de producție", "ore": 3, "audienta": "Cheie utilizatori",
   "grupe": [{"titlu": "Puncte de lucru",
              "subiecte": ["Definirea punctelor de lucru (Producție → Configurare → Puncte de lucru)"]}]},
  {"nume": "Raportarea producției", "ore": 2.5, "audienta": "Operatori",
   "grupe": [{"titlu": "Consumuri", "subiecte": ["Bon de consum suplimentar pentru pierderi 1-3%"]}]}
]}"""


async def test_modulele_vin_din_specificatie(tmp_path):
    with patch("anthropic.AsyncAnthropic", return_value=_raspuns(PAYLOAD_BUN)):
        module = await construieste_module_din_spec(_spec(tmp_path), "sk-x")

    assert [m["nume"] for m in module] == [
        "Configurări de producție",
        "Raportarea producției",
    ]
    assert [m["nr"] for m in module] == [1, 2]
    assert module[0]["ore"] == 3.0
    # calea din meniu, păstrată așa cum apare în document
    assert "Producție → Configurare" in module[0]["grupe"][0][1][0]


async def test_niciun_modul_protejat_de_comprimare(tmp_path):
    """Nimic nu e „standard", deci nimic nu e obligatoriu de păstrat."""
    with patch("anthropic.AsyncAnthropic", return_value=_raspuns(PAYLOAD_BUN)):
        module = await construieste_module_din_spec(_spec(tmp_path), "sk-x")
    assert all(m["esential"] is False for m in module)


async def test_specificatie_fara_material_de_training_da_eroare_explicita(tmp_path):
    with patch("anthropic.AsyncAnthropic", return_value=_raspuns('{"module": []}')):
        with pytest.raises(ValueError) as e:
            await construieste_module_din_spec(_spec(tmp_path), "sk-x")
    assert "niciun modul" in str(e.value)


async def test_document_gol_nu_ajunge_la_ai(tmp_path):
    from docx import Document

    p = tmp_path / "gol.docx"
    Document().save(str(p))
    client = _raspuns(PAYLOAD_BUN)
    with patch("anthropic.AsyncAnthropic", return_value=client):
        with pytest.raises(ValueError) as e:
            await construieste_module_din_spec(p, "sk-x")
    assert "text" in str(e.value)
    client.messages.create.assert_not_called()


async def test_raspuns_ciobit_se_repara(tmp_path):
    ciobit = '{"module": [{"nume": "Configurări", "ore": 2, "grupe": [{"titlu": "T", "subiecte": ["S"'
    with patch("anthropic.AsyncAnthropic", return_value=_raspuns(ciobit)):
        module = await construieste_module_din_spec(_spec(tmp_path), "sk-x")
    assert module[0]["nume"] == "Configurări"


def test_modul_fara_continut_este_ignorat():
    assert _curata_modul({"nume": "Gol", "ore": 2, "grupe": []}, 1) is None
    assert _curata_modul({"nume": "", "ore": 2}, 1) is None
    assert _curata_modul({"nume": "X", "grupe": [{"titlu": "T", "subiecte": []}]}, 1) is None


def test_orele_sunt_aduse_in_limite_rezonabile():
    def ore(valoare):
        m = _curata_modul(
            {"nume": "X", "ore": valoare, "grupe": [{"titlu": "T", "subiecte": ["S"]}]}, 1
        )
        return m["ore"]

    assert ore(2.3) == 2.5          # rotunjire la jumătăți de oră
    assert ore(0) == 0.5            # sub minim
    assert ore(40) == 8.0           # peste maxim
    assert ore("nu e număr") == 0.5


async def test_nu_se_accepta_mai_mult_de_limita_de_module(tmp_path):
    multe = ", ".join(
        f'{{"nume": "M{i}", "ore": 1, "grupe": [{{"titlu": "T", "subiecte": ["S"]}}]}}'
        for i in range(20)
    )
    payload = f'{{"module": [{multe}]}}'
    with patch("anthropic.AsyncAnthropic", return_value=_raspuns(payload)):
        module = await construieste_module_din_spec(_spec(tmp_path), "sk-x")
    assert len(module) == MAX_MODULE

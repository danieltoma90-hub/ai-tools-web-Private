# -*- coding: utf-8 -*-
"""Rutele de training si contractul lor de validare."""
import io
from unittest.mock import patch

from docx import Document

from auth import verify_token
from main import app


def _spec_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Specificatie Productie", level=1)
    doc.add_paragraph("Reteta se preia din sistemul de laborator.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _auth():
    app.dependency_overrides[verify_token] = lambda: {"id": "u1", "email": "t@x.ro"}


async def test_estimate_without_auth_returns_401(client):
    response = await client.get("/api/training/estimate?tip=core&zile=3")
    assert response.status_code == 401


async def test_estimate_core_3_zile(client):
    _auth()
    try:
        response = await client.get(
            "/api/training/estimate?tip=core&zile=3",
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 200
    data = response.json()
    assert data["buget_ore"] == 18.0
    assert data["ore_referinta"] == 17.0
    assert data["comprimat"] is False


async def test_estimate_o_zi_este_comprimat(client):
    _auth()
    try:
        response = await client.get(
            "/api/training/estimate?tip=core&zile=1",
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.json()["comprimat"] is True


async def test_tip_necunoscut_returns_422(client):
    _auth()
    try:
        response = await client.get(
            "/api/training/estimate?tip=logistica&zile=3",
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 422


async def test_zile_peste_limita_returns_422(client):
    _auth()
    try:
        response = await client.get(
            "/api/training/estimate?tip=core&zile=40",
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 422


async def test_productie_fara_specificatie_returns_422(client):
    _auth()
    try:
        response = await client.post(
            "/api/training/generate",
            data={"tip": "productie", "zile": "3", "client": "ACME"},
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 422
    assert "specificația" in response.json()["detail"].lower()


async def test_specificatie_in_alt_format_returns_422(client):
    _auth()
    try:
        response = await client.post(
            "/api/training/generate",
            data={"tip": "productie", "zile": "3", "client": "ACME"},
            files={"file": ("spec.pdf", b"%PDF-1.4", "application/pdf")},
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 422


async def test_core_fara_specificatie_porneste_job(client):
    """CORE merge si fara specificatie — genereaza programul standard."""
    _auth()
    try:
        response = await client.post(
            "/api/training/generate",
            data={"tip": "core", "zile": "3", "client": "ACME"},
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 200
    assert response.json()["job_id"]


async def test_specificatie_fara_cheie_api_returns_500(client, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    _auth()
    try:
        response = await client.post(
            "/api/training/generate",
            data={"tip": "productie", "zile": "3", "client": "ACME"},
            files={
                "file": (
                    "spec.docx",
                    _spec_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 500
    assert "ANTHROPIC_API_KEY" in response.json()["detail"]


async def test_specificatia_vine_din_storage(client, monkeypatch, tmp_path):
    """Traseul normal: fisierul urcat direct in Supabase, nu prin proxy."""
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    spec = tmp_path / "spec.docx"
    spec.write_bytes(_spec_bytes())
    _auth()
    try:
        with patch("routers.training.download_upload", return_value=spec) as descarcat:
            response = await client.post(
                "/api/training/generate",
                data={
                    "tip": "productie",
                    "zile": "3",
                    "client": "ACME",
                    "storage_path": "training/abc.docx",
                },
                headers={"Authorization": "Bearer fake"},
            )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 200
    descarcat.assert_called_once_with("training/abc.docx")


async def test_specificatie_disparuta_din_storage_returns_422(client, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    _auth()
    try:
        with patch("routers.training.download_upload", side_effect=Exception("not found")):
            response = await client.post(
                "/api/training/generate",
                data={
                    "tip": "productie",
                    "zile": "3",
                    "client": "ACME",
                    "storage_path": "training/lipsa.docx",
                },
                headers={"Authorization": "Bearer fake"},
            )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 422
    assert "Reîncarc" in response.json()["detail"]


async def test_job_inexistent_returns_404(client):
    _auth()
    try:
        response = await client.get(
            "/api/training/job/nu-exista",
            headers={"Authorization": "Bearer fake"},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 404

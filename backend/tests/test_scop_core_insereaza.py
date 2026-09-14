# -*- coding: utf-8 -*-
"""Teste pentru inserarea capitolului CORE cu elemente suplimentare.

`insereaza.py` (Task 1) a fost portat înghețat, înainte ca `capitol.construieste`
(Task 2) să capete parametrul `elemente` — de-atunci nu avea o cale să-l
primească pe al lui, deci a doua copie produsă de pipeline („gazdă cu
capitol”) conținea doar cele zece secțiuni standard, fără elementele
suplimentare ale clientului (limitare semnalată în raportul Task 4). Acest
fișier acoperă fix golul: `insereaza_capitol` primește acum `elemente` și le
duce mai departe la `capitol.construieste`, exact ca la capitolul standalone.

`GAZDA` (fixture-ul comun, vezi `test_scop_core_capitol.py`) nu conține în
forma ei brută niciun titlu de capitol — doar buline „List Paragraph” — deci
nu poate proba singură renumerotarea capitolelor care urmează după punctul de
inserare. `_gazda_cu_capitole_existente` construiește un fișier temporar
derivat din `GAZDA` (aceleași stiluri, deci `Heading 1` rezolvă real), cu
două capitole numerotate adăugate, ca renumerotarea să aibă ce muta.
"""
from __future__ import annotations

import hashlib
import pathlib

from docx import Document

from skills.scop_core import capitol, insereaza

GAZDA = pathlib.Path(__file__).resolve().parent / "fixtures" / "gazda_reala_anonimizata.docx"


def _hash(p: pathlib.Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _titluri(doc) -> list[tuple[str, str]]:
    """(nume_stil, text) pentru fiecare paragraf Heading 1/2/3, în ordine."""
    return [
        (p.style.name, p.text)
        for p in doc.paragraphs
        if p.style is not None and p.style.name in ("Heading 1", "Heading 2", "Heading 3")
    ]


def _gazda_cu_capitole_existente(tmp_path) -> pathlib.Path:
    """`GAZDA` + două capitole „Heading 1” numerotate, 5 și 6 — ca să existe
    ceva de renumerotat la inserarea unui capitol nou tot cu `numar=5`."""
    doc = Document(str(GAZDA))
    doc.add_paragraph("5. Capitolul vechi cinci", style="Heading 1")
    doc.add_paragraph("6. Capitolul vechi șase", style="Heading 1")
    cale = tmp_path / "gazda_cu_capitole.docx"
    doc.save(str(cale))
    return cale


# --- elementele suplimentare apar în documentul inserat ---------------------

def test_elementul_pe_sectiune_apare_ca_subcapitol_in_documentul_inserat():
    el = capitol.Element(titlu="Aprobare comenzi pe niveluri", text="Text.", plasare="achizitii")
    rezultat = insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME", elemente=[el])
    texte = [t for _, t in _titluri(rezultat)]
    # "achizitii" e a 6-a secțiune canonică (vezi ordinea din charisma_core.SECTIUNI).
    assert "5.6. Modulul Achiziții" in texte
    assert "5.6.1. Aprobare comenzi pe niveluri" in texte
    assert texte.index("5.6.1. Aprobare comenzi pe niveluri") > texte.index("5.6. Modulul Achiziții")


def test_elementul_propriu_apare_dupa_cele_zece_sectiuni_standard():
    el = capitol.Element(titlu="Integrare cu cântarul electronic", text="Text.", plasare="propriu")
    rezultat = insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME", elemente=[el])
    h2 = [t for stil_nume, t in _titluri(rezultat) if stil_nume == "Heading 2"]
    assert h2[-1] == "5.11. Integrare cu cântarul electronic"


# --- numerotarea rămâne corectă -----------------------------------------

def test_capitolele_gazdei_de_dupa_inserare_raman_renumerotate_cu_elemente_prezente(tmp_path):
    """Non-regresie: prezența elementelor suplimentare nu trebuie să strice
    renumerotarea capitolelor gazdei care urmează după punctul de inserare —
    logica aceea (`renumeroteaza`) e independentă de `elemente`, dar merită
    verificată explicit împreună, nu doar separat."""
    gazda = _gazda_cu_capitole_existente(tmp_path)
    el = capitol.Element(titlu="Element propriu", text="Text.", plasare="propriu")
    rezultat = insereaza.insereaza_capitol(gazda, capitol.alege_sectiuni(), client="ACME", elemente=[el])
    texte = [t for _, t in _titluri(rezultat)]
    assert "6. Capitolul vechi cinci" in texte
    assert "7. Capitolul vechi șase" in texte
    # capitolul nou (cu elementul lui propriu) apare înaintea celor renumerotate
    assert texte.index("5.11. Element propriu") < texte.index("6. Capitolul vechi cinci")


def test_subnumerele_elementelor_de_pe_aceeasi_sectiune_raman_contigue_in_documentul_inserat():
    e1 = capitol.Element(titlu="Primul", text="T1.", plasare="vanzari")
    e2 = capitol.Element(titlu="Al doilea", text="T2.", plasare="vanzari")
    rezultat = insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME", elemente=[e1, e2])
    texte = [t for _, t in _titluri(rezultat)]
    assert "5.5.1. Primul" in texte
    assert "5.5.2. Al doilea" in texte
    assert texte.index("5.5.1. Primul") < texte.index("5.5.2. Al doilea")


# --- compatibilitate cu apelanții existenți --------------------------------

def test_fara_elemente_insereaza_capitol_iese_exact_ca_inainte():
    """Apelanții existenți (care nu știu de `elemente`) trebuie să primească
    exact același document ca înainte de acest fix — omiterea parametrului,
    `elemente=None` explicit și `elemente=[]` explicit sunt echivalente."""
    fara_parametru = insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME")
    cu_none = insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME", elemente=None)
    cu_lista_goala = insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME", elemente=[])

    texte_fara_parametru = [p.text for p in fara_parametru.paragraphs]
    assert texte_fara_parametru == [p.text for p in cu_none.paragraphs]
    assert texte_fara_parametru == [p.text for p in cu_lista_goala.paragraphs]
    # nicio urmă de secțiune proprie suplimentară (5.11+) — doar cele zece standard
    h2 = [t for stil_nume, t in _titluri(fara_parametru) if stil_nume == "Heading 2"]
    assert h2[-1] == "5.10. Migrarea și inițializarea datelor"


# --- gazda de pe disc nu se modifică niciodată -----------------------------

def test_gazda_ramane_neschimbata_dupa_inserare_cu_elemente():
    """Non-negociabilul din brief, verificat explicit și pe calea cu elemente
    suplimentare — nu doar pe cea fără, deja acoperită în test_scop_core_pipeline.py."""
    hash_inainte = _hash(GAZDA)
    el = capitol.Element(titlu="X", text="Y.", plasare="propriu")
    insereaza.insereaza_capitol(GAZDA, capitol.alege_sectiuni(), client="ACME", elemente=[el])
    assert _hash(GAZDA) == hash_inainte

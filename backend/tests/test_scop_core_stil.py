# -*- coding: utf-8 -*-
"""Teste pentru `skills/scop_core/stil.py` — deschiderea curată a fișierelor
Word (item 4 din valul de fix-uri) și eliminarea imaginilor orfane rămase
după golirea corpului gazdei (item 7).

Verificarea completă a itemului 7 — mărimea reală înainte/după pe gazda
Turkish Doner Steakhouse, deschiderea documentului rezultat, randarea
titlurilor/bulinelor/borduri de tabel — trăiește în raportul manual de
verificare end-to-end, nu aici: acela e documentul real de la client, nu o
dependență a suitei de teste a acestui repo (vezi și docstring-ul
`test_scop_core_capitol.py` despre `GAZDA`). Aici se verifică mecanismul, cu
o imagine minimală, sintetică, adăugată programatic.
"""
from __future__ import annotations

import base64
import pathlib
import zipfile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Cm

from skills.scop_core import stil

GAZDA = pathlib.Path(__file__).resolve().parent / "fixtures" / "gazda_reala_anonimizata.docx"

# PNG minimal (1x1, transparent) — suficient ca python-docx să-l accepte ca
# imagine reală (are propriul parser de dimensiuni PNG, nu depinde de Pillow).
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY"
    "42YAAAAASUVORK5CYII="
)


def _gazda_cu_imagine(tmp_path) -> pathlib.Path:
    """`GAZDA` + o imagine inline în corp — ca să poată fi verificată
    eliminarea ei ca relație orfană după golirea corpului (item 7)."""
    doc = Document(str(GAZDA))
    png = tmp_path / "mini.png"
    png.write_bytes(_PNG_1X1)
    doc.add_picture(str(png))
    cale = tmp_path / "gazda_cu_imagine.docx"
    doc.save(str(cale))
    return cale


# --- deschide_docx / DocumentInvalid (item 4) --------------------------------

def test_deschide_docx_pe_fisier_invalid_da_documentinvalid_curat(tmp_path):
    """Un fișier care nu e deloc zip (deci nici Word) trebuie să dea
    `DocumentInvalid`, cu un mesaj care NU conține calea temporară a
    fișierului — spre deosebire de excepția brută a lui python-docx."""
    cale = tmp_path / "nu_e_docx.docx"
    cale.write_bytes(b"nu sunt un docx valid deloc")

    with pytest.raises(stil.DocumentInvalid) as exc_info:
        stil.deschide_docx(cale)

    mesaj = str(exc_info.value)
    assert str(cale) not in mesaj
    assert cale.name not in mesaj


def test_document_din_gazda_pe_fisier_invalid_da_documentinvalid(tmp_path):
    """`document_din_gazda` trece prin `deschide_docx` — trebuie să propage
    aceeași excepție curată, nu `ValueError`-ul brut al python-docx."""
    cale = tmp_path / "nu_e_docx.docx"
    cale.write_bytes(b"nici asta nu e un docx")

    with pytest.raises(stil.DocumentInvalid):
        stil.document_din_gazda(cale)


# --- eliminarea imaginilor orfane (item 7) -----------------------------------

def test_document_din_gazda_elimina_imaginea_orfana_din_corp(tmp_path):
    """Corpul e golit de `document_din_gazda`, dar imaginea rămasă în
    `word/media/` (fostul conținut al gazdei) nu mai e referită de nimic —
    trebuie eliminată, nu cărată mai departe în capitolul generat."""
    gazda_cu_imagine = _gazda_cu_imagine(tmp_path)
    with zipfile.ZipFile(gazda_cu_imagine) as z:
        assert any(n.startswith("word/media/") for n in z.namelist()), (
            "fixture-ul de test trebuie să conțină efectiv o imagine înainte de fix"
        )

    doc = stil.document_din_gazda(gazda_cu_imagine)

    rels_imagine = [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]
    assert rels_imagine == []

    iesire = tmp_path / "capitol_fara_imagine.docx"
    doc.save(str(iesire))
    with zipfile.ZipFile(iesire) as z:
        parti_media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert parti_media == [], f"imaginea orfană tot ajunge în fișier: {parti_media}"
    assert iesire.stat().st_size < gazda_cu_imagine.stat().st_size


def test_document_din_gazda_fara_nicio_imagine_nu_pica(tmp_path):
    """Non-regresie: gazda obișnuită (fără nicio imagine în corp, cazul
    `GAZDA` normal) trebuie să treacă prin `_elimina_imagini_orfane` fără
    nicio eroare — lista de relații de tip imagine e goală, nu lipsă."""
    doc = stil.document_din_gazda(GAZDA)
    assert [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE] == []


# --- antet/subsol: detectare mismatch de client și golire ------------------
# GAZDA are antetul/subsolul golite (vezi docstring-ul fișierului) — cu un
# singur paragraf gol fiecare. Testele de mai jos pornesc de la o COPIE a
# gazdei, cu text real adăugat în antet/subsol, ca gazda reală (Turkish Doner
# Steakhouse, unde numele clientului stă într-un TABEL din subsol, nu într-un
# paragraf simplu) — vezi `_gazda_cu_antet_subsol`.

def _gazda_cu_antet_subsol(tmp_path, antet_text: str = "", subsol_tabel: list | None = None) -> pathlib.Path:
    doc = Document(str(GAZDA))
    s = doc.sections[0]
    if antet_text:
        s.header.paragraphs[0].add_run(antet_text)
    if subsol_tabel:
        t = s.footer.add_table(rows=1, cols=len(subsol_tabel), width=Cm(16))
        for celula, text in zip(t.rows[0].cells, subsol_tabel):
            celula.text = text
    cale = tmp_path / "gazda_cu_antet_subsol.docx"
    doc.save(str(cale))
    return cale


def test_texte_antet_subsol_pe_gazda_goala_intoarce_siruri_goale():
    doc = stil.document_din_gazda(GAZDA)
    antet, subsol = stil.texte_antet_subsol(doc)
    assert antet == ""
    assert subsol == ""


def test_texte_antet_subsol_citeste_text_din_paragraf(tmp_path):
    gazda = _gazda_cu_antet_subsol(tmp_path, antet_text="Descrierea soluției | Modul Producție")
    doc = stil.document_din_gazda(gazda)
    antet, subsol = stil.texte_antet_subsol(doc)
    assert "Descrierea soluției" in antet
    assert subsol == ""


def test_texte_antet_subsol_citeste_text_din_tabelul_subsolului(tmp_path):
    """Gazda reală ține numele clientului într-un tabel din subsol
    („Turkish Doner Steakhouse S.R.L. | Charisma ERP | v1.0”), nu într-un
    paragraf simplu — `texte_antet_subsol` trebuie să-l găsească acolo."""
    gazda = _gazda_cu_antet_subsol(
        tmp_path, subsol_tabel=["", "Turkish Doner Steakhouse S.R.L. | Charisma ERP | v1.0", "Pag. 1 / 1"],
    )
    doc = stil.document_din_gazda(gazda)
    _, subsol = stil.texte_antet_subsol(doc)
    assert "Turkish Doner Steakhouse S.R.L." in subsol


def test_curata_antet_subsol_goleste_paragraful_antetului(tmp_path):
    gazda = _gazda_cu_antet_subsol(tmp_path, antet_text="Vechi Client SRL")
    doc = stil.document_din_gazda(gazda)
    stil.curata_antet_subsol(doc)
    antet, _ = stil.texte_antet_subsol(doc)
    assert antet == ""


def test_curata_antet_subsol_goleste_tabelul_subsolului(tmp_path):
    gazda = _gazda_cu_antet_subsol(
        tmp_path, subsol_tabel=["", "Vechi Client SRL | Charisma ERP | v1.0", "Pag. 1 / 1"],
    )
    doc = stil.document_din_gazda(gazda)
    stil.curata_antet_subsol(doc)
    _, subsol = stil.texte_antet_subsol(doc)
    assert subsol == ""


def test_curata_antet_subsol_nu_atinge_un_run_doar_cu_imagine(tmp_path):
    """Un logo în antet stă într-un run fără text vizibil (`run.text == ""`,
    doar un `w:drawing`) — golirea trebuie să-l lase intact, nu să-l șteargă
    ca efect secundar al lui `CT_R.clear_content` pe un run fără text."""
    doc = Document(str(GAZDA))
    header = doc.sections[0].header
    png = tmp_path / "logo.png"
    png.write_bytes(_PNG_1X1)
    run_logo = header.paragraphs[0].add_run()
    run_logo.add_picture(str(png))
    header.paragraphs[0].add_run("Vechi Client SRL")
    cale = tmp_path / "gazda_cu_logo.docx"
    doc.save(str(cale))

    doc2 = stil.document_din_gazda(cale)
    rels_imagine_inainte = [r for r in doc2.part.rels.values() if r.reltype == RT.IMAGE]
    stil.curata_antet_subsol(doc2)

    antet, _ = stil.texte_antet_subsol(doc2)
    assert antet == ""  # textul a fost golit

    iesire = tmp_path / "rezultat.docx"
    doc2.save(str(iesire))
    with zipfile.ZipFile(iesire) as z:
        parti_media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert parti_media, "logo-ul din antet nu trebuie șters de golirea textului"


def test_curata_antet_subsol_pastreaza_pachetul_deschis_curat(tmp_path):
    """Verificarea „lecției” din brief: golirea antetului/subsolului nu
    trebuie să lase relații orfane sau un pachet pe care Word să-l respingă —
    fișierul rezultat tot trebuie să se deschidă curat, cu `sectPr` și
    referințele lui de antet/subsol intacte."""
    gazda = _gazda_cu_antet_subsol(
        tmp_path, antet_text="Vechi Client SRL",
        subsol_tabel=["", "Vechi Client SRL | Charisma ERP | v1.0", "Pag. 1 / 1"],
    )
    doc = stil.document_din_gazda(gazda)
    stil.curata_antet_subsol(doc)
    iesire = tmp_path / "rezultat.docx"
    doc.save(str(iesire))

    # se redeschide fără nicio excepție — un pachet stricat (relații orfane,
    # părți lipsă) ar ridica una aici.
    redeschis = Document(str(iesire))
    sectPr = redeschis.sections[0]._sectPr
    assert sectPr.find(qn("w:headerReference")) is not None
    assert sectPr.find(qn("w:footerReference")) is not None

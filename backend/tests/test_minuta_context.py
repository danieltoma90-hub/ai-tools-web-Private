# -*- coding: utf-8 -*-
"""Teste pentru contextul de proiect al minutei."""
import sys
from pathlib import Path

import pytest
from docx import Document

SKILL_SCRIPTS = Path(__file__).parent.parent / "skills" / "minuta" / "scripts"
sys.path.insert(0, str(SKILL_SCRIPTS))

from build_context_template import build_context_template  # noqa: E402
from context_parser import parse_context  # noqa: E402
from pipelines.minuta_pipeline import _followup_section  # noqa: E402


@pytest.fixture
def template(tmp_path) -> Path:
    return build_context_template(tmp_path / "template.docx")


def test_template_has_all_chapters(template):
    doc = Document(str(template))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(heads) == 7
    assert any("Participanți" in h for h in heads)
    assert any("Glosar" in h for h in heads)
    assert any("Acțiuni deschise" in h for h in heads)


def test_empty_template_yields_empty_context(template):
    """Un formular necompletat NU trebuie să ajungă în prompt ca informație."""
    ctx = parse_context(template)
    assert ctx.is_empty()
    assert ctx.as_prompt_block() == ""


def _fill(template: Path, out: Path) -> Path:
    doc = Document(str(template))
    for i, val in enumerate(["Godac", "Charisma CORE", "GOD-1", "Analiza", "Daniel"], start=1):
        doc.tables[0].rows[i].cells[1].text = val
    doc.tables[1].rows[1].cells[0].text = "Ana Popescu"
    doc.tables[1].rows[1].cells[1].text = "Director economic"
    doc.tables[1].rows[1].cells[2].text = "Godac"
    doc.tables[2].rows[1].cells[0].text = "NIR"
    doc.tables[2].rows[1].cells[1].text = "Notă de intrare-recepție"
    doc.tables[3].rows[1].cells[0].text = "12.06.2026"
    doc.tables[3].rows[1].cells[1].text = "Se folosește FIFO"
    doc.tables[4].rows[1].cells[0].text = "Godac"
    doc.tables[4].rows[1].cells[1].text = "Trimite structura gestiunilor"
    doc.tables[4].rows[1].cells[2].text = "01.07.2026"
    doc.save(str(out))
    return out


def test_filled_context_parsed_and_rendered(template, tmp_path):
    ctx = parse_context(_fill(template, tmp_path / "filled.docx"))
    assert not ctx.is_empty()
    assert ctx.proiect["Client"] == "Godac"
    assert ctx.participanti[0]["rol"] == "Director economic"
    assert ctx.glosar[0]["termen"] == "NIR"
    assert ctx.decizii[0]["decizie"] == "Se folosește FIFO"
    assert ctx.actiuni_deschise[0]["termen"] == "01.07.2026"

    block = ctx.as_prompt_block()
    assert "Ana Popescu" in block and "Director economic" in block
    assert "NIR" in block
    # instrucțiunile pentru model, nu doar datele
    assert "NU le raporta ca noi" in block
    assert "urmareste in transcript" in block


def test_hints_and_labels_are_not_content(template, tmp_path):
    """Instrucțiunile italice și etichetele („Module / arii în scop:") nu sunt răspunsuri."""
    ctx = parse_context(_fill(template, tmp_path / "f2.docx"))
    joined = " ".join(ctx.scop + ctx.preferinte)
    assert "Exemple:" not in joined
    assert "Module / arii în scop:" not in joined
    assert "Completați" not in joined


def test_followup_section_built_from_model_output():
    section = _followup_section({"stadiu_actiuni": [
        {"actiune": "Trimite gestiuni", "responsabil": "Godac",
         "stare": "Finalizată", "detaliu": "Transmis pe 3 sept."},
        {"actiune": "Mediu de test", "responsabil": "TotalSoft", "stare": "Nediscutată"},
    ]})
    assert section["titlu"] == "Stadiul acțiunilor anterioare"
    rows = section["blocuri"][0]["rows"]
    assert len(rows) == 2
    assert rows[0][2] == "Finalizată"
    assert rows[1][3] == "—"  # fără detaliu, nu string gol


def test_followup_section_absent_when_nothing_to_report():
    assert _followup_section(None) is None
    assert _followup_section({"stadiu_actiuni": []}) is None
    # intrări fără acțiune sunt ignorate, nu produc rânduri goale
    assert _followup_section({"stadiu_actiuni": [{"stare": "Finalizată"}]}) is None

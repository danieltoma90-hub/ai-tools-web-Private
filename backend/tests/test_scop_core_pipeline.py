# -*- coding: utf-8 -*-
"""Teste pentru pipeline-ul scop_core (Task 4 — orchestrare capitol + inserare).

`GAZDA` e același fixture folosit de `test_scop_core_capitol.py` — un document
gazdă real, anonimizat, cu styles.xml/numbering.xml reale. Testul central de
siguranță (`test_gazda_ramane_neschimbata_dupa_inserare`) verifică prin hash
că fișierul de pe disc nu e niciodată atins, indiferent de calea de execuție —
vezi non-negociabilul din brief.
"""
from __future__ import annotations

import hashlib
import pathlib

import pytest
from docx import Document

from pipelines import scop_core_pipeline as pipeline
from skills.scop_core import extractie
from skills.scop_core.charisma_core import SECTIUNI

GAZDA = pathlib.Path(__file__).resolve().parent / "fixtures" / "gazda_reala_anonimizata.docx"


def _hash(p: pathlib.Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _texte(path: pathlib.Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


def _h2(path: pathlib.Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs if p.style is not None and p.style.name == "Heading 2"]


@pytest.fixture
def spion_mktemp(monkeypatch):
    """Înlocuiește `_mktemp_path` cu o variantă care reține fiecare cale creată,
    ca testele de eșec să poată verifica direct — nu doar presupune — că
    niciun fișier temporar nu rămâne pe disc după o excepție."""
    create: list[pathlib.Path] = []
    original = pipeline._mktemp_path

    def spion(suffix):
        p = original(suffix)
        create.append(p)
        return p

    monkeypatch.setattr(pipeline, "_mktemp_path", spion)
    return create


# --- capitolul se produce -----------------------------------------------

async def test_pipelineul_produce_capitolul_cu_numele_clientului():
    capitol_path, gazda_out, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME SRL", elemente=[], insereaza_in_gazda=False,
    )
    try:
        assert capitol_path.is_file()
        text = "\n".join(_texte(capitol_path))
        assert "Charisma ERP CORE" in text
        assert "ACME SRL" in text
        assert gazda_out is None
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_client_gol_foloseste_placeholder_implicit():
    capitol_path, _, _ = await pipeline.run_scop_core_pipeline(
        GAZDA, client="", elemente=[], insereaza_in_gazda=False,
    )
    try:
        text = "\n".join(_texte(capitol_path))
        assert pipeline.CLIENT_IMPLICIT in text
    finally:
        capitol_path.unlink(missing_ok=True)


# --- a doua cale (gazdă cu capitol inserat) ------------------------------

async def test_insereaza_in_gazda_true_produce_a_doua_cale():
    capitol_path, gazda_out, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=True,
    )
    try:
        assert gazda_out is not None
        assert gazda_out.is_file()
        assert gazda_out != capitol_path
        assert sumar["gazda_inserata"] is True
    finally:
        capitol_path.unlink(missing_ok=True)
        if gazda_out is not None:
            gazda_out.unlink(missing_ok=True)


async def test_insereaza_in_gazda_false_intoarce_none():
    capitol_path, gazda_out, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=False,
    )
    try:
        assert gazda_out is None
        assert sumar["gazda_inserata"] is False
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_gazda_ramane_neschimbata_dupa_inserare():
    """Non-negociabilul din brief: fișierul gazdă de pe disc nu se modifică
    niciodată — hash identic înainte/după, chiar cu insereaza_in_gazda=True."""
    hash_inainte = _hash(GAZDA)
    capitol_path, gazda_out, _ = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=True,
    )
    try:
        assert _hash(GAZDA) == hash_inainte
    finally:
        capitol_path.unlink(missing_ok=True)
        if gazda_out is not None:
            gazda_out.unlink(missing_ok=True)


# --- elemente malformate din browser --------------------------------------

async def test_element_fara_titlu_e_respins_dar_nu_pica_pipeline_ul():
    elemente = [
        {"titlu": "", "text": "Text valid.", "plasare": "vanzari", "fluxuri_legate": []},
        {"titlu": "Titlu bun", "text": "Text bun.", "plasare": "vanzari", "fluxuri_legate": []},
    ]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["elemente_respinse"] == 1
        assert sumar["elemente_plasate"] == 1
        assert "Titlu bun" in "\n".join(_texte(capitol_path))
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_element_fara_text_e_respins():
    elemente = [{"titlu": "Titlu fără corp", "text": "   ", "plasare": "propriu", "fluxuri_legate": []}]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["elemente_respinse"] == 1
        assert sumar["elemente_plasate"] == 0
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_element_care_nu_e_dict_e_ignorat():
    elemente = ["nu sunt un dict", 42, {"titlu": "T", "text": "X", "plasare": "propriu"}]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["elemente_respinse"] == 2
        assert sumar["elemente_plasate"] == 1
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_plasare_invalida_devine_propriu_si_nu_pica():
    elemente = [{"titlu": "Element rătăcit", "text": "Y", "plasare": "sectiune-care-nu-exista",
                 "fluxuri_legate": []}]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["elemente_proprii"] == 1
        assert sumar["elemente_pe_sectiune"] == 0
        assert _h2(capitol_path)[-1].endswith("Element rătăcit")
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_plasare_lipsa_devine_propriu():
    elemente = [{"titlu": "X", "text": "Y", "fluxuri_legate": []}]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["elemente_proprii"] == 1
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_elemente_care_nu_e_lista_nu_pica_pipeline_ul():
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=None, insereaza_in_gazda=False,  # type: ignore[arg-type]
    )
    try:
        assert sumar["elemente_plasate"] == 0
        assert sumar["elemente_respinse"] == 0
        assert sumar["elemente_primite"] == 0
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_fluxuri_legate_care_nu_e_lista_devine_lista_goala():
    elemente = [{"titlu": "X", "text": "Y", "plasare": "vanzari", "fluxuri_legate": "V2"}]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["elemente_plasate"] == 1
        assert "corelat" not in "\n".join(_texte(capitol_path)).lower()
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_cod_de_flux_valid_produce_fraza_de_corelare():
    elemente = [{"titlu": "X", "text": "Y", "plasare": "vanzari", "fluxuri_legate": [" V2 "]}]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        v2 = next(f for s in SECTIUNI if s.cheie == "vanzari" for f in s.fluxuri if f.cod == "V2")
        assert v2.flux in "\n".join(_texte(capitol_path))
    finally:
        capitol_path.unlink(missing_ok=True)


# --- sumarul ---------------------------------------------------------------

async def test_sumarul_are_numerele_corecte():
    elemente = [
        {"titlu": "Sub A", "text": "Text A.", "plasare": "vanzari", "fluxuri_legate": []},
        {"titlu": "Sub B", "text": "Text B.", "plasare": "propriu", "fluxuri_legate": []},
    ]
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=False,
    )
    try:
        assert sumar["sectiuni"] == len(SECTIUNI) == 10
        assert sumar["fluxuri"] == sum(len(s.fluxuri) for s in SECTIUNI) == 47
        assert sumar["elemente_primite"] == 2
        assert sumar["elemente_plasate"] == 2
        assert sumar["elemente_pe_sectiune"] == 1
        assert sumar["elemente_proprii"] == 1
        assert sumar["elemente_respinse"] == 0
    finally:
        capitol_path.unlink(missing_ok=True)


# --- on_step ----------------------------------------------------------------

async def test_on_step_semnaleaza_parsing_si_building_fara_inserare():
    pasi: list[str] = []
    capitol_path, _, _ = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=False, on_step=pasi.append,
    )
    try:
        assert pasi[0] == "parsing"
        assert "building" in pasi
        assert "inserare" not in pasi
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_on_step_semnaleaza_si_inserarea_cand_e_ceruta():
    pasi: list[str] = []
    capitol_path, gazda_out, _ = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=True, on_step=pasi.append,
    )
    try:
        assert "inserare" in pasi
        assert pasi.index("building") < pasi.index("inserare")
    finally:
        capitol_path.unlink(missing_ok=True)
        if gazda_out is not None:
            gazda_out.unlink(missing_ok=True)


# --- avertisment -------------------------------------------------------------

async def test_avertisment_gol_fara_elemente_si_fara_insereaza():
    capitol_path, _, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=False,
    )
    try:
        assert sumar["avertisment"] == ""
    finally:
        capitol_path.unlink(missing_ok=True)


async def test_elementele_suplimentare_ajung_si_in_gazda_cu_capitol():
    """Limitarea din raportul Task 4 a fost rezolvată: `insereaza.insereaza_capitol`
    primește acum `elemente` — cele două fișiere întoarse de pipeline conțin
    identic aceleași elemente suplimentare, nu doar capitolul standalone.
    Fără avertisment: nu mai există nimic de semnalat pe această cale."""
    elemente = [{"titlu": "Element unic pentru test", "text": "Y", "plasare": "propriu",
                 "fluxuri_legate": []}]
    capitol_path, gazda_out, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=elemente, insereaza_in_gazda=True,
    )
    try:
        assert sumar["avertisment"] == ""
        assert gazda_out is not None
        assert "Element unic pentru test" in "\n".join(_texte(gazda_out))
        assert "5.11. Element unic pentru test" in _h2(gazda_out)
    finally:
        capitol_path.unlink(missing_ok=True)
        if gazda_out is not None:
            gazda_out.unlink(missing_ok=True)


async def test_fara_elemente_insereaza_in_gazda_nu_are_avertisment():
    capitol_path, gazda_out, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=True,
    )
    try:
        assert sumar["avertisment"] == ""
    finally:
        capitol_path.unlink(missing_ok=True)
        if gazda_out is not None:
            gazda_out.unlink(missing_ok=True)


# --- eșecuri și curățare de fișiere temporare -------------------------------

async def test_gazda_inexistenta_ridica_eroare_fara_fisier_temporar_ramas(tmp_path, spion_mktemp):
    gazda_inexistenta = tmp_path / "nu_exista.docx"
    with pytest.raises(FileNotFoundError):
        await pipeline.run_scop_core_pipeline(
            gazda_inexistenta, client="ACME", elemente=[], insereaza_in_gazda=False,
        )
    assert spion_mktemp, "cel puțin un fișier temporar trebuia creat înainte de eroare"
    assert not any(p.exists() for p in spion_mktemp), "niciun fișier temporar nu trebuie să rămână pe disc"


async def test_esecul_la_inserare_pastreaza_capitolul_si_avertizeaza(monkeypatch, spion_mktemp):
    """Un eșec la pasul opțional de inserare nu trebuie să tragă după el
    capitolul deja construit cu succes — vezi docstring-ul funcției."""
    def insereaza_stricata(cale_gazda, sectiuni, numar=5, nivel=1, client="[NUME CLIENT]",
                           elemente=None):
        raise RuntimeError("gazdă simulat coruptă")

    monkeypatch.setattr(pipeline.insereaza, "insereaza_capitol", insereaza_stricata)

    capitol_path, gazda_out, sumar = await pipeline.run_scop_core_pipeline(
        GAZDA, client="ACME", elemente=[], insereaza_in_gazda=True,
    )
    try:
        assert capitol_path.is_file()
        assert gazda_out is None
        assert sumar["gazda_inserata"] is False
        assert "eșuat" in sumar["avertisment"]
        # candidatul creat pentru a doua cale nu trebuie să rămână orfan pe disc
        candidati_ramasi = [p for p in spion_mktemp if p != capitol_path and p.exists()]
        assert candidati_ramasi == []
    finally:
        capitol_path.unlink(missing_ok=True)


# --- propune_job -------------------------------------------------------------

async def test_propune_job_e_wrapper_subtire_peste_extractie(monkeypatch, tmp_path):
    primit: dict = {}

    async def fals(path):
        primit["path"] = path
        return [{"titlu": "X", "text": "Y", "plasare": "propriu", "fluxuri_legate": []}]

    monkeypatch.setattr(pipeline.extractie, "propune_elemente", fals)

    supliment = tmp_path / "supliment.docx"
    supliment.touch()
    rezultat = await pipeline.propune_job(supliment)

    assert primit["path"] == supliment
    assert rezultat == [{"titlu": "X", "text": "Y", "plasare": "propriu", "fluxuri_legate": []}]


async def test_propune_job_propaga_erorile_lui_extractie(monkeypatch, tmp_path):
    async def fals(path):
        raise ValueError("Răspunsul modelului nu este JSON valid")

    monkeypatch.setattr(pipeline.extractie, "propune_elemente", fals)

    supliment = tmp_path / "supliment.docx"
    supliment.touch()
    with pytest.raises(ValueError, match="JSON valid"):
        await pipeline.propune_job(supliment)

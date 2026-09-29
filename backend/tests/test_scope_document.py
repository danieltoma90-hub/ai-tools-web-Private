# -*- coding: utf-8 -*-
"""Teste pentru documentul complet de scop CORE (Task 2 din planul scop-core-web).

`scope.py` e portul lui `d:\\AI_Claude\\skills\\generare-scope-core\\scope.py`:
asamblează cele 11 capitole ale documentului de ofertă (nu doar capitolul CORE
standard, ca `capitol.py`). Testele de aici verifică comportamentele
documentate ca fiind obligatoriu păstrate la port — numerotarea care se
strânge fără goluri, capitolul de delimitare condiționat de documentul-frate,
erorile pentru elemente orfane sau ordine incompletă, forma capitolului de
acoperire, transliterarea și rescrierea antetului — nu reimplementează suita
completă a skill-ului local (aceea rămâne sursa de adevăr, în celălalt repo).

`GAZDA` e fixture-ul comun al pachetului `scop_core` (vezi
`test_scop_core_capitol.py`) — folosit ca document-gazdă, ca testele să nu
depindă de niciun fișier din `d:\\AI_Claude`.
"""
from __future__ import annotations

import copy
import pathlib
import re

import docx
import pytest

from skills.scop_core import scope
from skills.scop_core.charisma_core import Flux

GAZDA = pathlib.Path(__file__).resolve().parent / "fixtures" / "gazda_reala_anonimizata.docx"

CONFIG: dict = {
    "client": {
        "nume": "Client Test SRL",
        "nume_complet": "Client Test SRL",
        "domeniu": "activitate de test",
        "entitati": [],
        "observatii": "",
    },
    "document": {
        "titlu": "Descrierea soluției ofertate",
        "subtitlu": "Implementare Charisma ERP CORE",
        "versiune": "1.0",
        "elaborat_de": "Echipa de consultanță TotalSoft, Delivery ERP",
        "furnizor": "TotalSoft S.A.",
    },
    "document_frate": {"exista": False, "titlu": "", "arie_acoperita": "", "coduri": []},
    "sectiuni_core": {"toate": True, "doar": [], "fara": []},
    "capitole": {
        "context": True, "abordare": True, "beneficii": True, "acoperire": True,
        "delimitare": True, "premise": True, "confirmari": True, "sinteza": True,
        "validare": True,
    },
    "stil": {"document_gazda": str(GAZDA), "antet": "Antet de test"},
}


def _genereaza(tmp_path, cfg=None, **kw):
    cfg = copy.deepcopy(cfg or CONFIG)
    cale = scope.genereaza(cfg, tmp_path / "rezultat.docx", **kw)
    return docx.Document(str(cale))


def _titluri_h1(doc) -> list[str]:
    """Doar titlurile de capitol (Heading 1), în ordinea din document."""
    return [p.text.strip() for p in doc.paragraphs
            if p.style.name == "Heading 1" and p.text.strip()]


def _text(doc) -> str:
    parti = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parti.extend(c.text for c in r.cells)
    return "\n".join(parti)


# --- 1. Toate capitolele activate → 11 capitole, în ordine -------------------

def test_toate_capitolele_activate_produce_11_capitole_in_ordine(tmp_path):
    """Cu toate flag-urile din `capitole` pe True și `document_frate.exista`
    pe True (singura condiție suplimentară pentru capitolul de delimitare),
    documentul are exact cele 11 capitole standard, numerotate 1..11 în
    ordinea din brief."""
    cfg = copy.deepcopy(CONFIG)
    cfg["document_frate"] = {"exista": True, "titlu": "Scop Producție",
                             "arie_acoperita": "Producția", "coduri": ["N1 - N6"]}
    titluri = _titluri_h1(_genereaza(tmp_path, cfg))
    numere = [int(re.match(r"^(\d+)\.", t).group(1)) for t in titluri]
    assert numere == list(range(1, 12)), f"așteptam 11 capitole 1..11, am obținut {numere}"

    asteptate = [
        "Scopul documentului", "Contextul proiectului",
        "Abordarea și metodologia de implementare",
        "Soluția ofertată — Charisma ERP CORE", "Beneficiile soluției",
        "Analiza de acoperire a cerințelor", "Delimitarea față de documentul-frate",
        "Premise și responsabilități", "Puncte de confirmat înainte de validarea scopului",
        "Sinteza scopului funcțional", "Validarea scopului",
    ]
    for titlu, parte in zip(titluri, asteptate):
        assert parte in titlu, f"capitol neașteptat: {titlu!r} (așteptam ...{parte!r})"


# --- 2. Dezactivarea capitolelor strânge numerotarea, fără goluri -----------

def test_dezactivarea_capitolelor_stringe_numerotarea_fara_goluri(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg["capitole"]["context"] = False
    cfg["capitole"]["abordare"] = False
    cfg["capitole"]["confirmari"] = False
    doc = _genereaza(tmp_path, cfg)
    titluri = _titluri_h1(doc)
    numere = [int(re.match(r"^(\d+)\.", t).group(1)) for t in titluri]
    assert numere == list(range(1, len(numere) + 1)), f"numerotare discontinuă: {numere}"
    # capitolele dezactivate chiar lipsesc, nu doar renumerotate sub alt număr
    corp = " | ".join(titluri)
    assert "Contextul proiectului" not in corp
    assert "Abordarea și metodologia" not in corp
    assert "Puncte de confirmat" not in corp
    # iar «Soluția ofertată» — care e mereu scris — ajunge acum al doilea capitol
    assert "Soluția ofertată" in titluri[1]


# --- 3. Capitolul de delimitare apare doar cu document_frate.exista --------

def test_delimitarea_apare_doar_cu_document_frate_exista(tmp_path):
    fara_frate = _titluri_h1(_genereaza(tmp_path))
    assert not any("Delimitarea" in t for t in fara_frate)

    cfg = copy.deepcopy(CONFIG)
    cfg["document_frate"] = {"exista": True, "titlu": "Scop Producție",
                             "arie_acoperita": "Producția", "coduri": ["N1 - N6"]}
    cu_frate = _titluri_h1(_genereaza(tmp_path, cfg))
    assert any("Delimitarea" in t for t in cu_frate)


def test_delimitarea_nu_apare_daca_document_frate_exista_e_fals_desi_capitolul_e_activat(tmp_path):
    """`capitole.delimitare = True` nu ajunge — fără `document_frate.exista`
    capitolul tot nu se scrie (regula e o conjuncție, nu un singur steag)."""
    cfg = copy.deepcopy(CONFIG)
    cfg["capitole"]["delimitare"] = True
    cfg["document_frate"]["exista"] = False
    titluri = _titluri_h1(_genereaza(tmp_path, cfg))
    assert not any("Delimitarea" in t for t in titluri)


# --- 4. Un Suplimentar atașat unui modul exclus ridică eroare ---------------

def test_suplimentar_atasat_unui_modul_exclus_ridica_eroare(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg["sectiuni_core"] = {"toate": False, "doar": ["contabilitate"], "fara": []}
    orfan = scope.Suplimentar(titlu="Element orfan", puncte=["Punct."], in_modul="financiar")
    with pytest.raises(ValueError, match="financiar"):
        _genereaza(tmp_path, cfg, suplimentare=[orfan])


# --- 5. ORDINE_CAP4 incompletă sau cu element necunoscut ridică eroare -----

def test_ordine_cap4_incompleta_ridica_eroare(tmp_path):
    """Ordinea listează doar o parte din secțiunile efectiv incluse — lipsește
    «financiar» — și trebuie semnalat, ca să nu dispară tăcut din document."""
    cfg = copy.deepcopy(CONFIG)
    cfg["sectiuni_core"] = {"toate": False, "doar": ["contabilitate", "financiar"], "fara": []}
    with pytest.raises(ValueError, match="financiar"):
        _genereaza(tmp_path, cfg, ordine_cap4=["contabilitate"])


def test_ordine_cap4_cu_element_necunoscut_ridica_eroare(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg["sectiuni_core"] = {"toate": False, "doar": ["contabilitate", "financiar"], "fara": []}
    with pytest.raises(ValueError, match="nu-exista-asa-ceva"):
        _genereaza(tmp_path, cfg, ordine_cap4=["contabilitate", "financiar", "nu-exista-asa-ceva"])


def test_ordine_cap4_completa_reordoneaza_subcapitolele(tmp_path):
    """O ordine completă și validă chiar produce efectul cerut — secțiunea
    listată a doua apare înaintea celei listate întâi în configurarea implicită."""
    cfg = copy.deepcopy(CONFIG)
    cfg["sectiuni_core"] = {"toate": False, "doar": ["contabilitate", "financiar"], "fara": []}
    doc = _genereaza(tmp_path, cfg, ordine_cap4=["financiar", "contabilitate"])
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    idx_financiar = next(i for i, t in enumerate(h2) if "Modulul Financiar" in t)
    idx_contabilitate = next(i for i, t in enumerate(h2) if "Modulul Contabilitate" in t)
    assert idx_financiar < idx_contabilitate


# --- forma_acoperire necunoscută ridică eroare; formele cunoscute schimbă forma -

def test_forma_acoperire_necunoscuta_ridica_eroare(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg.setdefault("stil", {})["forma_acoperire"] = "nu-exista"
    with pytest.raises(ValueError, match="[Ff]orm"):
        _genereaza(tmp_path, cfg, acoperire=[scope.Cerinta("Z", "X", "R.", "A")])


def test_forma_acoperire_apdn_produce_tabelul_de_incadrari(tmp_path):
    doc = _genereaza(tmp_path)  # implicit "apdn"
    titluri = _titluri_h1(doc)
    assert any("Analiza de acoperire a cerințelor" in t for t in titluri)
    assert "Modul de încadrare a cerințelor" in _text(doc)


def test_forma_acoperire_trei_grupe_schimba_titlul_si_gruparea(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg.setdefault("stil", {})["forma_acoperire"] = "trei-grupe"
    cerinte = [
        scope.Cerinta("Z", "Acoperit", "R.", "A"),
        scope.Cerinta("Z", "Parțial", "R.", "P"),
        scope.Cerinta("Z", "De definit", "R.", "D"),
        scope.Cerinta("Z", "În afară", "R.", "N"),
    ]
    titluri = _titluri_h1(_genereaza(tmp_path, cfg, acoperire=cerinte))
    assert any("Scopul ofertat și delimitări" in t for t in titluri)
    assert "Modul de încadrare a cerințelor" not in _text(_genereaza(tmp_path, cfg, acoperire=cerinte))


# --- fara_diacritice transliterează efectiv documentul ----------------------

DIACRITICE = "ăĂâÂîÎșȘşŞțȚţŢ"


def test_fara_diacritice_transliterreaza_intregul_document(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg.setdefault("stil", {})["fara_diacritice"] = True
    cfg["stil"]["antet"] = "Antet cu ăîșțâ"
    doc = _genereaza(tmp_path, cfg, acoperire=[scope.Cerinta("Zonă", "Cerință", "Răspuns.", "A")])

    text = _text(doc)
    for p in doc.sections[0].header.paragraphs:
        text += "\n" + p.text
    ramase = sorted({c for c in text if c in DIACRITICE})
    assert not ramase, f"diacritice rămase în document: {ramase}"
    assert "Solutia ofertata" in text


def test_diacriticele_raman_implicit_cand_steagul_lipseste(tmp_path):
    text = _text(_genereaza(tmp_path))
    assert any(c in DIACRITICE for c in text)


# --- antetul se rescrie mereu din configurare (documentul-gazdă e frate) ----

def test_antetul_se_rescrie_din_configurare(tmp_path):
    """Gazda folosită la test e chiar fixture-ul comun — dar regula generală
    e că antetul documentului final vine din `stil.antet`, nu din gazdă:
    dacă gazda e documentul-frate, antetul ei numește aria altui document."""
    doc = _genereaza(tmp_path)
    antete = [p.text.strip() for p in doc.sections[0].header.paragraphs if p.text.strip()]
    assert antete == ["Antet de test"], f"antet neașteptat: {antete}"


# --- documentul nu conține niciodată «perimetru» ----------------------------

def test_documentul_nu_contine_cuvantul_perimetru(tmp_path):
    cfg = copy.deepcopy(CONFIG)
    cfg["document_frate"] = {"exista": True, "titlu": "Scop Producție",
                             "arie_acoperita": "Producția", "coduri": ["N1 - N6"]}
    sup = scope.Suplimentar(titlu="Integrare de test", intro="Intro.", puncte=["Punct unu."],
                            fluxuri=[Flux(cod="EX1", flux="Mapare", presupune="Ecran de mapare.")])
    cerinte = [scope.Cerinta("Z", "Cerință", "Răspuns.", "A")]
    doc = _genereaza(tmp_path, cfg, suplimentare=[sup], acoperire=cerinte)
    assert "perimetr" not in _text(doc).lower()


# --- ACOPERIRE grupează pe A/P/D/N și sare peste categoriile goale ----------

def test_acoperirea_grupeaza_cerintele_pe_incadrari(tmp_path):
    cerinte = [
        scope.Cerinta("Zona A", "Cerință acoperită", "Răspuns.", "A"),
        scope.Cerinta("Zona B", "Cerință parțială", "Răspuns.", "P"),
        scope.Cerinta("Zona C", "Cerință de definit", "Răspuns.", "D"),
        scope.Cerinta("Zona D", "Cerință în afară", "Răspuns.", "N"),
    ]
    doc = _genereaza(tmp_path, acoperire=cerinte)
    titluri = [p.text.strip() for p in doc.paragraphs
               if p.style.name.startswith("Heading") and p.text.strip()]
    for eticheta in ("(A)", "(P)", "(D)", "(N)"):
        assert any(eticheta in t for t in titluri), f"lipsește subcapitolul {eticheta}"
    text = _text(doc)
    for c in cerinte:
        assert c.cerinta in text
    assert "Au fost analizate 4 elemente" in text


def test_acoperirea_sare_peste_incadrarile_fara_cerinte(tmp_path):
    doc = _genereaza(tmp_path, acoperire=[scope.Cerinta("Z", "Doar A", "Răspuns.", "A")])
    titluri = [p.text.strip() for p in doc.paragraphs
               if p.style.name.startswith("Heading") and p.text.strip()]
    assert any("(A)" in t for t in titluri)
    assert not any("(P)" in t or "(D)" in t or "(N)" in t for t in titluri)

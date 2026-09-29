# -*- coding: utf-8 -*-
"""Pipeline scop_core — orchestrarea skill-ului `skills/scop_core` pentru web.

Fluxul din UI are doi pași separați:
  1. Utilizatorul încarcă un document suplimentar (opțional) -> `propune_job`
     îl trimite la model (prin `extractie.propune_elemente`) și întoarce
     elementele propuse, ca să le corecteze în browser.
  2. Utilizatorul confirmă (eventual editate) elementele și cere generarea ->
     `run_scop_core_pipeline` construiește capitolul standard CORE, cu
     elementele suplimentare deja plasate, și — dacă i se cere — și o copie
     a documentului gazdă cu capitolul inserat la locul lui.

Elementele care ajung la `run_scop_core_pipeline` sunt dicturi simple
(`{"titlu", "text", "plasare", "fluxuri_legate"}`) care au trecut printr-un
formular din browser și prin editările utilizatorului — nu se poate avea
încredere oarbă în formă. Conversia lor în `capitol.Element` se face AICI,
nu în router (routerul nu importă skill-ul direct), cu aceeași filosofie de
validare „nu ridica, corectează sau aruncă” pe care `extractie._valideaza_elemente`
o aplică deja părții AI a fluxului — vezi `_element_din_dict`.

`insereaza.insereaza_capitol` primește acum aceleași `elemente` cu care se
construiește și capitolul separat — cele două fișiere întoarse de acest
pipeline (capitolul standalone și, dacă e cerută, gazda cu capitolul inserat)
conțin identic aceleași elemente suplimentare, plasate la fel.

`delimitari` (intrările „Delimitări de scop”) urmează exact același drum ca
`elemente`: dicturi brute din browser, validate tolerant (`_delimitare_din_dict`)
și duse identic pe ambele căi. `curata_antet_subsol` golește antetul/subsolul
moștenite din gazdă pe ambele căi, pentru cazul unei gazde reutilizate ca
șablon de stil pentru alt client — vezi `stil.curata_antet_subsol`.

A doua orchestrare din acest fișier, `run_scope_document_pipeline`, construiește
DOCUMENTUL COMPLET de scop (11 capitole, `skills.scop_core.scope`) — nu doar
capitolul CORE de mai sus. Primește `continut`, un dicționar brut din browser
cu toate elementele editabile (client, elemente suplimentare, matricea de
acoperire A/P/D/N, delimitarea față de documentul-frate, fluxul operațional,
ordinea capitolului 4, beneficii, puncte de confirmat) și le convertește
defensiv în dataclass-urile din `scope.py`, cu aceeași filosofie „nu ridica,
corectează sau numără” ca mai sus — vezi `_client_din_dict`,
`_suplimentare_din_dicturi`, `_acoperire_din_dicturi` etc. Excepție de la acea
filosofie: erorile pe care `scope.genereaza` le ridică deliberat (element
suplimentar atașat unui modul exclus; ordinea capitolului 4 incompletă sau cu
element necunoscut) NU se înghit aici — se lasă să treacă neschimbate, ca
routerul să le transforme într-un 422 cu mesaj clar, nu într-o pierdere tăcută
de conținut dintr-o ofertă comercială.

Marcajul `scope.PLACEHOLDER` ("[ de completat ]") e comportamentul corect al
lui `scope.genereaza` pentru skill-ul local, unde omul completează manual
golurile direct în Word — dar pe web nu există acel pas manual, iar
documentul rezultat e o ofertă comercială care poate ajunge direct la un
client plătitor. Un capitol pe care ecranul nu are cum să-l umple cu conținut
real trebuie OPRIT (comutatorul lui din `capitole` rămâne pe `False`), nu
completat cu marcaje — dar câteva goluri (situația actuală/obiectivele
clientului, tabelul de delimitare fără rânduri) nu au comutator propriu și
tot pot ajunge scrise cu placeholder dacă formularul nu le-a completat.
`_gaseste_placeholdere`, mai jos, e plasa de siguranță de după construcție:
citește documentul deja scris, paragraf cu paragraf ȘI celulă de tabel cu
celulă de tabel, și numără fiecare apariție rămasă a marcajului — rezultatul
ajunge în `sumar` (`placeholder_numar`, `placeholder_capitole`) și în
`avertisment`, ca utilizatorul să vadă înainte să trimită fișierul mai
departe, nu să descopere ulterior că a trimis un `[ de completat ]` unui
client.
"""
from __future__ import annotations

import copy
import logging
import os
import tempfile
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from skills.scop_core import capitol, extractie, insereaza, scope, stil
from skills.scop_core.charisma_core import SECTIUNI, Flux
from skills.scop_core.stil import DocumentInvalid  # re-exportat pentru router,
# care nu importă `skills.scop_core` direct — vezi docstring-ul de sus.

logger = logging.getLogger(__name__)

# cele zece chei de secțiune CORE, plus "propriu" — singurele valori valide
# pentru `plasare`, aceeași listă pe care se bazează și `extractie.py`.
_PLASARI_VALIDE: set[str] = {s.cheie for s in SECTIUNI} | {"propriu"}

CLIENT_IMPLICIT = "[NUME CLIENT]"


def _mktemp_path(suffix: str) -> Path:
    fd, name = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    return Path(name)


def valideaza_docx(cale: Path) -> None:
    """Verifică sincron că `cale` chiar se deschide ca document Word.

    Nu păstrează rezultatul — e o gardă rapidă, apelată de router imediat
    după descărcarea gazdei în `/scop-core/genereaza`, ÎNAINTE de a porni
    jobul de fundal. Fără ea, un fișier cu extensia `.docx` dar conținut
    invalid (ex. un `.xlsx` redenumit) trecea de verificarea de extensie,
    pornea un job care eșua abia mai târziu în `stil.document_din_gazda`, iar
    utilizatorul primea un job „error" cu mesaj tehnic în loc de un răspuns
    422 imediat — vezi raportul, itemul despre erorile de upload malformat.
    Ridică `DocumentInvalid` (mesaj curat, fără cale de fișier) dacă `cale`
    nu e un `.docx` valid.
    """
    stil.deschide_docx(cale)


async def propune_job(supliment_path: Path) -> list[dict]:
    """Subțire wrapper peste `extractie.propune_elemente`.

    Există doar ca routerul să nu importe `skills.scop_core` direct — la fel
    cum fac celelalte perechi router/pipeline din acest repo (vezi
    `mockup_pipeline`/`scenarii_pipeline`). Nu adaugă nicio logică proprie:
    validarea elementelor propuse de model trăiește deja în `extractie.py`,
    iar erorile ei (JSON invalid, cheie API lipsă) urcă neschimbate.
    """
    return await extractie.propune_elemente(supliment_path)


def _element_din_dict(brut: object) -> capitol.Element | None:
    """Convertește un singur dict brut (din browser) într-un `capitol.Element`.

    Regulile — în oglindă cu `extractie._valideaza_elemente`, ca un element
    editat de utilizator și unul propus direct de model să fie tratate la
    fel de tolerant:
      - orice altceva decât un `dict` (None, șir, listă imbricată etc.) e
        respins — nu are cum să fie un element valid;
      - `titlu`/`text` lipsă, goale după strip sau de alt tip decât `str`
        fac elementul inutilizabil (nu poți randa un titlu de capitol gol,
        nici un sub-capitol fără conținut) — respins, nu „reparat” cu text
        de rezervă inventat aici;
      - `plasare` care nu e exact una din cele zece chei de secțiune CORE
        sau `"propriu"` devine tăcut `"propriu"` — elementul nu se pierde,
        doar capătă secțiune proprie la coada capitolului, exact cum
        `capitol.construieste` tratează deja o secțiune filtrată afară
        (vezi docstring-ul funcției respective);
      - `fluxuri_legate` care nu e o listă (sau lipsește) devine `[]`; intrările
        care nu sunt șiruri se aruncă, iar spațiile parazite se curăță cu
        `.strip()` — codurile invalide propriu-zise (inexistente în
        `SECTIUNI`) NU se filtrează aici, ci mai departe, în
        `capitol._denumiri_fluxuri_legate`, care e sursa unică a acestei
        reguli — nu se duplică logica ei aici.
    """
    if not isinstance(brut, dict):
        return None

    titlu_brut = brut.get("titlu")
    titlu = titlu_brut.strip() if isinstance(titlu_brut, str) else ""
    text_brut = brut.get("text")
    text = text_brut.strip() if isinstance(text_brut, str) else ""
    if not titlu or not text:
        return None

    plasare = brut.get("plasare")
    if not isinstance(plasare, str) or plasare not in _PLASARI_VALIDE:
        plasare = "propriu"

    fluxuri_brut = brut.get("fluxuri_legate")
    fluxuri_legate = [
        cod.strip() for cod in fluxuri_brut if isinstance(cod, str) and cod.strip()
    ] if isinstance(fluxuri_brut, list) else []

    return capitol.Element(titlu=titlu, text=text, plasare=plasare, fluxuri_legate=fluxuri_legate)


def _elemente_din_dicturi(brute: object) -> tuple[list[capitol.Element], int]:
    """Convertește lista de dicturi brute; întoarce (elemente valide, nr. respinse).

    O `elemente` care nu e deloc listă (None, obiect greșit trimis) e tratată
    ca listă goală — fără nicio excepție ridicată către router.
    """
    if not isinstance(brute, list):
        return [], 0

    valide: list[capitol.Element] = []
    respinse = 0
    for brut in brute:
        element = _element_din_dict(brut)
        if element is None:
            respinse += 1
        else:
            valide.append(element)
    return valide, respinse


def _delimitare_din_dict(brut: object) -> capitol.Delimitare | None:
    """Convertește un singur dict brut (din browser) într-un `capitol.Delimitare`.

    Aceeași filosofie tolerantă ca `_element_din_dict`: doar `dict`-uri sunt
    acceptate, iar `element`/`precizare` lipsă, goale după strip sau de alt
    tip decât `str` fac intrarea inutilizabilă — un rând din tabelul
    „Element | Precizare” fără unul din cele două câmpuri nu are ce însemna,
    și nu se completează tăcut cu text de rezervă inventat aici.
    """
    if not isinstance(brut, dict):
        return None

    element_brut = brut.get("element")
    element = element_brut.strip() if isinstance(element_brut, str) else ""
    precizare_brut = brut.get("precizare")
    precizare = precizare_brut.strip() if isinstance(precizare_brut, str) else ""
    if not element or not precizare:
        return None

    return capitol.Delimitare(element=element, precizare=precizare)


def _delimitari_din_dicturi(brute: object) -> tuple[list[capitol.Delimitare], int]:
    """Convertește lista de dicturi brute de delimitări; întoarce (valide, respinse).

    O `delimitari` care nu e deloc listă (None, obiect greșit trimis) e
    tratată ca listă goală — fără nicio excepție ridicată către router,
    exact ca la `_elemente_din_dicturi`.
    """
    if not isinstance(brute, list):
        return [], 0

    valide: list[capitol.Delimitare] = []
    respinse = 0
    for brut in brute:
        d = _delimitare_din_dict(brut)
        if d is None:
            respinse += 1
        else:
            valide.append(d)
    return valide, respinse


async def run_scop_core_pipeline(
    gazda_path: Path,
    client: str,
    elemente: list[dict],
    insereaza_in_gazda: bool,
    delimitari: list[dict] | None = None,
    curata_antet_subsol: bool = False,
    on_step=None,
) -> tuple[Path, Path | None, dict]:
    """Construiește capitolul CORE și, opțional, copia gazdă+capitol.

    Întoarce `(capitol_path, gazda_cu_capitol_path, sumar)`. A doua cale e
    `None` dacă `insereaza_in_gazda` e `False`, SAU dacă a fost cerută dar
    pasul de inserare a eșuat — vezi mai jos: acel eșec nu duce la pierderea
    capitolului deja construit cu succes, ci doar la un `avertisment` în sumar
    (aceeași filosofie ca în `training_pipeline._module_core`: un pas
    secundar care pică nu trebuie să tragă după el livrabilul principal).

    Documentul gazdă NU e niciodată scris pe disc: atât `capitol.construieste`
    cât și `insereaza.insereaza_capitol` lucrează pe un `Document` încărcat în
    memorie (`stil.document_din_gazda`/`Document(gazda_path)` intern) — abia
    rezultatul se salvează, sub o cale nouă, temporară.

    `delimitari` — intrările „Delimitări de scop” (dicturi brute din browser,
    convertite tolerant prin `_delimitari_din_dicturi`, în oglindă cu
    `elemente`), duse la fel spre `capitol.construieste`/`insereaza_capitol`
    și randate ca ultimul sub-capitol pe ambele căi.

    `curata_antet_subsol` — golește antetul/subsolul moștenite din gazdă (vezi
    `stil.curata_antet_subsol`), pentru cazul unei gazde reutilizate ca șablon
    de stil pentru alt client. Implicit `False`. Când e `False` ȘI clientul a
    fost completat efectiv (nu placeholder-ul implicit), sumarul semnalează
    dacă numele lui nu apare deloc în antetul/subsolul găsit — vezi
    `antet_subsol_avertisment` mai jos.
    """
    if on_step:
        on_step("parsing")

    sectiuni = capitol.alege_sectiuni()
    client_declarat = (client or "").strip()
    client_nume = client_declarat or CLIENT_IMPLICIT
    elemente_valide, elemente_respinse = _elemente_din_dicturi(elemente)
    delimitari_valide, delimitari_respinse = _delimitari_din_dicturi(delimitari)

    chei_sectiuni = {s.cheie for s in sectiuni}
    elemente_pe_sectiune = sum(1 for el in elemente_valide if el.plasare in chei_sectiuni)
    elemente_proprii = len(elemente_valide) - elemente_pe_sectiune

    if on_step:
        on_step("building")

    capitol_path = _mktemp_path(".docx")
    try:
        doc = stil.document_din_gazda(gazda_path)
        antet_text, subsol_text = stil.texte_antet_subsol(doc)
        if curata_antet_subsol:
            stil.curata_antet_subsol(doc)
        capitol.construieste(doc, sectiuni, client=client_nume, elemente=elemente_valide,
                             delimitari=delimitari_valide)
        doc.save(str(capitol_path))
    except Exception:
        capitol_path.unlink(missing_ok=True)
        raise

    # Mismatch de client în antetul/subsolul moștenit din gazdă — doar dacă
    # utilizatorul chiar a completat un client (altfel comparăm cu placeholder-ul
    # implicit, care evident n-apare niciodată) ȘI n-a cerut deja curățarea lor
    # (caz în care mismatch-ul e deja rezolvat, nu mai e nimic de semnalat).
    antet_subsol_avertisment = ""
    if not curata_antet_subsol and client_declarat and (antet_text or subsol_text):
        text_gasit = f"{antet_text}\n{subsol_text}".lower()
        if client_declarat.lower() not in text_gasit:
            bucati = []
            if subsol_text:
                bucati.append(f"subsol: „{subsol_text}”")
            if antet_text:
                bucati.append(f"antet: „{antet_text}”")
            antet_subsol_avertisment = (
                f"Numele clientului („{client_declarat}”) nu apare în antetul/subsolul moștenit "
                "din documentul gazdă — verifică dacă documentul a fost creat pentru alt client "
                "înainte să-l trimiți. Text găsit — " + "; ".join(bucati) + "."
            )

    gazda_path_out: Path | None = None
    avertismente: list[str] = []
    if elemente_respinse:
        avertismente.append(
            f"{elemente_respinse} element(e) suplimentar(e) primite din browser nu au putut fi "
            "folosite (titlu sau text lipsă) și au fost ignorate."
        )
    if delimitari_respinse:
        avertismente.append(
            f"{delimitari_respinse} intrare(intrări) de „Delimitări de scop” primite din browser "
            "nu au putut fi folosite (element sau precizare lipsă) și au fost ignorate."
        )

    if insereaza_in_gazda:
        if on_step:
            on_step("inserare")
        candidat = _mktemp_path(".docx")
        try:
            gazda_cu_capitol = insereaza.insereaza_capitol(
                gazda_path, sectiuni, client=client_nume, elemente=elemente_valide,
                delimitari=delimitari_valide, curata_antet_subsol=curata_antet_subsol,
            )
            gazda_cu_capitol.save(str(candidat))
            gazda_path_out = candidat
        except DocumentInvalid as e:
            # Mesajul lui `DocumentInvalid` e garantat curat (fără cale de
            # fișier) — sigur de arătat direct utilizatorului, spre deosebire
            # de orice altă excepție de mai jos.
            candidat.unlink(missing_ok=True)
            avertismente.append(f"Inserarea capitolului în documentul gazdă a eșuat: {e}")
        except Exception as e:
            # Excepție neprevăzută: nu interpolăm `{e}` brut într-un mesaj
            # către utilizator (poate conține o cale de fișier de pe server
            # sau alt detaliu tehnic) — detaliul rămâne în logul serverului,
            # utilizatorul primește un mesaj generic.
            candidat.unlink(missing_ok=True)
            logger.warning("Inserarea capitolului în gazdă a eșuat: %s", e)
            avertismente.append("Inserarea capitolului în documentul gazdă a eșuat.")

    sumar = {
        "sectiuni": len(sectiuni),
        "fluxuri": sum(len(s.fluxuri) for s in sectiuni),
        "elemente_primite": len(elemente) if isinstance(elemente, list) else 0,
        "elemente_plasate": len(elemente_valide),
        "elemente_pe_sectiune": elemente_pe_sectiune,
        "elemente_proprii": elemente_proprii,
        "elemente_respinse": elemente_respinse,
        "delimitari_primite": len(delimitari) if isinstance(delimitari, list) else 0,
        "delimitari_plasate": len(delimitari_valide),
        "delimitari_respinse": delimitari_respinse,
        "gazda_inserata": gazda_path_out is not None,
        "antet_subsol_curatat": curata_antet_subsol,
        "antet_subsol_avertisment": antet_subsol_avertisment,
        "avertisment": " ".join(avertismente),
    }
    return capitol_path, gazda_path_out, sumar


# =============================================================================
# Documentul complet de scop (11 capitole) — Task 3 din planul scop-core-web.
# =============================================================================

_INCADRARI_VALIDE: set[str] = {"A", "P", "D", "N"}


def _str_sau_gol(valoare: object) -> str:
    """`valoare.strip()` dacă e un `str` nevid după strip, altfel șir gol."""
    return valoare.strip() if isinstance(valoare, str) else ""


def _lista_stringuri(brute: object) -> list[str]:
    """Filtrează o listă brută, păstrând doar șirurile nevide (curățate de spații).

    O `brute` care nu e deloc listă (None, dict trimis din greșeală) devine
    listă goală — fără nicio excepție ridicată."""
    if not isinstance(brute, list):
        return []
    return [s.strip() for s in brute if isinstance(s, str) and s.strip()]


def _situatie_actuala_din_lista(brute: object) -> list[tuple[str, str]]:
    """Rândurile «Situația actuală | Cum se adresează în Charisma» ale clientului.

    Fiecare rând e un dict `{"actual": ..., "solutie": ...}`; un rând căruia
    îi lipsește oricare din cele două texte se ignoră tăcut — nu e unul din
    cele două erori care trebuie lăsate să treacă (vezi docstring-ul de sus),
    doar un rând de context necompletat, nu conținut de ofertă pierdut."""
    if not isinstance(brute, list):
        return []
    randuri = []
    for rand in brute:
        if not isinstance(rand, dict):
            continue
        actual = _str_sau_gol(rand.get("actual"))
        solutie = _str_sau_gol(rand.get("solutie"))
        if actual and solutie:
            randuri.append((actual, solutie))
    return randuri


def _client_din_dict(brut: object) -> scope.Client:
    """Construiește `scope.Client` dintr-un dict brut din browser.

    Câmpurile lipsă sau de tip greșit cad pe valorile implicite ale
    dataclass-ului (`scope.PLACEHOLDER` pentru `nume`, șiruri/liste goale
    pentru rest) — niciodată nu ridică excepție, ca un formular parțial
    completat de utilizator să nu blocheze restul generării."""
    if not isinstance(brut, dict):
        brut = {}
    kwargs: dict = {}
    nume = _str_sau_gol(brut.get("nume"))
    if nume:
        kwargs["nume"] = nume
    kwargs["nume_complet"] = _str_sau_gol(brut.get("nume_complet"))
    kwargs["domeniu"] = _str_sau_gol(brut.get("domeniu"))
    kwargs["entitati"] = _lista_stringuri(brut.get("entitati"))
    kwargs["situatie_actuala"] = _situatie_actuala_din_lista(brut.get("situatie_actuala"))
    kwargs["obiective"] = _lista_stringuri(brut.get("obiective"))
    kwargs["observatii"] = _str_sau_gol(brut.get("observatii"))
    return scope.Client(**kwargs)


def _fluxuri_din_lista(brute: object) -> list[Flux]:
    """`Flux`-urile unui element suplimentar — cod, denumire, ce presupune.

    Toate cele trei câmpuri sunt obligatorii pentru ca un rând de tabel să
    aibă sens; un flux căruia îi lipsește unul se ignoră (nu e conținut
    esențial al ofertei, doar detalierea unui element care rămâne oricum
    scris, chiar fără acel rând de tabel)."""
    if not isinstance(brute, list):
        return []
    rezultat = []
    for f in brute:
        if not isinstance(f, dict):
            continue
        cod = _str_sau_gol(f.get("cod"))
        flux_text = _str_sau_gol(f.get("flux"))
        presupune = _str_sau_gol(f.get("presupune"))
        if cod and flux_text and presupune:
            rezultat.append(Flux(cod=cod, flux=flux_text, presupune=presupune))
    return rezultat


def _suplimentar_din_dict(brut: object) -> scope.Suplimentar | None:
    """Un element suplimentar (capitol propriu sau atașat unui modul CORE).

    `titlu` lipsă/gol face elementul inutilizabil — respins, nu reparat cu
    text de rezervă. `in_modul`, dacă e prezent, se păstrează AȘA CUM VINE,
    fără validare împotriva secțiunilor CORE existente aici — acea validare
    e responsabilitatea lui `scope.genereaza` (`_cap_solutie`), care ridică
    `ValueError` pentru un modul exclus; ridicarea aceea NU trebuie duplicată
    sau anticipată aici, ca mesajul și punctul unic de adevăr să rămână
    `scope.py` — vezi docstring-ul de sus al fișierului."""
    if not isinstance(brut, dict):
        return None
    titlu = _str_sau_gol(brut.get("titlu"))
    if not titlu:
        return None
    in_modul_brut = brut.get("in_modul")
    in_modul = in_modul_brut.strip() if isinstance(in_modul_brut, str) and in_modul_brut.strip() else None
    incadrare_brut = brut.get("incadrare")
    incadrare = incadrare_brut.strip() if isinstance(incadrare_brut, str) and incadrare_brut.strip() else "Inclus"
    return scope.Suplimentar(
        titlu=titlu,
        intro=_str_sau_gol(brut.get("intro")),
        puncte=_lista_stringuri(brut.get("puncte")),
        fluxuri=_fluxuri_din_lista(brut.get("fluxuri")),
        nota=_str_sau_gol(brut.get("nota")),
        in_modul=in_modul,
        incadrare=incadrare,
    )


def _suplimentare_din_dicturi(brute: object) -> tuple[list[scope.Suplimentar], int]:
    """Convertește lista de elemente suplimentare; întoarce (valide, respinse)."""
    if not isinstance(brute, list):
        return [], 0
    valide: list[scope.Suplimentar] = []
    respinse = 0
    for brut in brute:
        element = _suplimentar_din_dict(brut)
        if element is None:
            respinse += 1
        else:
            valide.append(element)
    return valide, respinse


def _cerinta_din_dict(brut: object) -> tuple[scope.Cerinta | None, bool]:
    """O cerință a matricei de acoperire; întoarce `(cerința, a_fost_corectată)`.

    `cerinta`/`raspuns` lipsă fac rândul inutilizabil — respins. `zona` poate
    lipsi (rândul tot are sens, doar coloana «Zonă» rămâne goală). `incadrare`
    necunoscută sau lipsă NU se respinge — s-ar pierde o cerință reală doar
    pentru o etichetă greșită — ci se trece la «D. De definit», încadrarea
    cea mai conservatoare (cere analiză explicită, nu o declară nici acoperită,
    nici în afara scopului), și se semnalează prin `a_fost_corectată`."""
    if not isinstance(brut, dict):
        return None, False
    cerinta_text = _str_sau_gol(brut.get("cerinta"))
    raspuns = _str_sau_gol(brut.get("raspuns"))
    if not cerinta_text or not raspuns:
        return None, False
    zona = _str_sau_gol(brut.get("zona"))
    incadrare_brut = brut.get("incadrare")
    incadrare = incadrare_brut.strip().upper() if isinstance(incadrare_brut, str) else ""
    corectata = incadrare not in _INCADRARI_VALIDE
    if corectata:
        incadrare = "D"
    return scope.Cerinta(zona=zona, cerinta=cerinta_text, raspuns=raspuns, incadrare=incadrare), corectata


def _acoperire_din_dicturi(brute: object) -> tuple[list[scope.Cerinta], int, int]:
    """Convertește matricea de acoperire; întoarce (valide, respinse, corectate)."""
    if not isinstance(brute, list):
        return [], 0, 0
    valide: list[scope.Cerinta] = []
    respinse = 0
    corectate = 0
    for brut in brute:
        cerinta, corectata = _cerinta_din_dict(brut)
        if cerinta is None:
            respinse += 1
        else:
            valide.append(cerinta)
            if corectata:
                corectate += 1
    return valide, respinse, corectate


def _randuri_3col_din_dicturi(brute: object, chei: tuple[str, str, str]) -> tuple[list[list[str]], int]:
    """Rândurile unui tabel cu 3 coloane text (delimitare, flux operațional).

    Fiecare rând e un dict cu exact cheile din `chei`; un rând căruia îi
    lipsește oricare dintre cele trei valori se respinge — un rând pe jumătate
    completat dintr-un tabel contractual (delimitare) sau operațional (flux)
    n-are cum să se randeze coerent cu text de rezervă inventat aici."""
    if not isinstance(brute, list):
        return [], 0
    valide: list[list[str]] = []
    respinse = 0
    for brut in brute:
        if not isinstance(brut, dict):
            respinse += 1
            continue
        valori = [_str_sau_gol(brut.get(cheie)) for cheie in chei]
        if all(valori):
            valide.append(valori)
        else:
            respinse += 1
    return valide, respinse


def _confirmare_din_dict(brut: object) -> tuple[str, str] | None:
    if not isinstance(brut, dict):
        return None
    aspect = _str_sau_gol(brut.get("aspect"))
    motiv = _str_sau_gol(brut.get("motiv"))
    if not aspect or not motiv:
        return None
    return (aspect, motiv)


def _confirmari_din_dicturi(brute: object) -> tuple[list[tuple[str, str]], int]:
    if not isinstance(brute, list):
        return [], 0
    valide: list[tuple[str, str]] = []
    respinse = 0
    for brut in brute:
        confirmare = _confirmare_din_dict(brut)
        if confirmare is None:
            respinse += 1
        else:
            valide.append(confirmare)
    return valide, respinse


def _ordine_cap4_din_lista(brut: object) -> list[str] | None:
    """Ordinea subcapitolelor capitolului 4 — listă de chei de secțiune CORE
    amestecate cu titluri de elemente suplimentare cu capitol propriu.

    Intrările care nu sunt șiruri (sau devin goale după strip) se elimină
    tăcut AICI — dar dacă eliminarea lor face ordinea incompletă sau lasă
    o cheie necunoscută, `scope.genereaza` ridică `ValueError`, care NU se
    prinde în funcția asta: e exact eroarea deliberată care trebuie să
    ajungă la router, nu o pierdere tăcută de conținut. `None` (nu listă,
    sau listă golită complet de filtrare) înseamnă „nicio ordine cerută” —
    `scope.genereaza` foloseşte atunci ordinea canonică implicită."""
    if not isinstance(brut, list):
        return None
    curatata = [x.strip() for x in brut if isinstance(x, str) and x.strip()]
    return curatata or None


def _cfg_pregatit(cfg: object, gazda_path: Path | None) -> dict:
    """O copie sigură a `cfg` pentru `scope.genereaza`.

    `cfg` vine din browser, la fel ca `continut` — sub-dicturile structurale
    (`document`, `document_frate`, `sectiuni_core`, `capitole`, `stil`) devin
    goale dacă nu sunt chiar dicturi, ca accesul `.get(...)` din `scope.py` să
    nu ridice `AttributeError` pe un payload malformat.

    `stil.document_gazda` NU se preia NICIODATĂ din `cfg`: un șir venit din
    browser ar fi o cale de fișier aleasă de client, citită direct de pe
    discul serverului de `stil.document_din_gazda` — o cale de a deschide
    orice `.docx` accesibil procesului, nu doar documentele utilizatorului.
    Singura gazdă acceptată e cea descărcată de router din storage-ul
    utilizatorului (`gazda_path`, un fișier temporar local), sau nicio gazdă."""
    cfg = cfg if isinstance(cfg, dict) else {}
    rezultat = copy.deepcopy(cfg)

    for cheie in ("document", "document_frate", "sectiuni_core", "capitole", "stil"):
        if not isinstance(rezultat.get(cheie), dict):
            rezultat[cheie] = {}

    coduri = rezultat["document_frate"].get("coduri")
    rezultat["document_frate"]["coduri"] = list(coduri) if isinstance(coduri, list) else []

    for cheie in ("fara", "doar"):
        valoare = rezultat["sectiuni_core"].get(cheie)
        rezultat["sectiuni_core"][cheie] = (
            [c for c in valoare if isinstance(c, str)] if isinstance(valoare, list) else []
        )

    rezultat["stil"]["document_gazda"] = str(gazda_path) if gazda_path else None
    return rezultat


def _blocuri_document(doc):
    """Iterează paragrafele și tabelele documentului, în ordinea din corp (`w:body`).

    `doc.paragraphs` și `doc.tables` (python-docx) sunt liste plate și
    separate — nu păstrează ordinea relativă dintre ele, deci din ele nu se
    poate afla sub ce titlu de capitol se află un anumit tabel. Aici se
    citește direct elementul XML al corpului, care chiar păstrează acea
    ordine, și se învelește fiecare copil relevant (paragraf sau tabel) în
    tipul lui python-docx corespunzător — restul copiilor (ex. `w:sectPr`,
    proprietățile secțiunii) se ignoră, n-au text de căutat în ei."""
    for copil in doc.element.body.iterchildren():
        if copil.tag == qn("w:p"):
            yield Paragraph(copil, doc)
        elif copil.tag == qn("w:tbl"):
            yield Table(copil, doc)


def _gaseste_placeholdere(doc) -> tuple[int, list[str]]:
    """Caută marcajul `scope.PLACEHOLDER` care a ajuns totuși în documentul construit.

    Plasa de siguranță din docstring-ul de sus al fișierului: rulează DUPĂ ce
    `scope.genereaza` a scris deja documentul și numără fiecare apariție
    rămasă a marcajului, atât în paragrafe (situația actuală, obiectivele,
    beneficiile, acoperirea, punctele de confirmat pot ajunge scrise cu el)
    cât și în celule de tabel (delimitarea față de documentul-frate, sinteza)
    — un tabel NU apare deloc în `doc.paragraphs`, de-aici nevoia de
    `_blocuri_document` mai sus, care le vede pe amândouă în ordine.

    Întoarce `(numar_total, capitole)`. `capitole` e lista titlurilor de
    capitol/subcapitol (cel mai recent „Heading 1” SAU „Heading 2” văzut până
    la acel punct — orice conține markerul, la orice nivel din cele două) sub
    care a apărut cel puțin un marker, în ordinea întâlnirii, fiecare o
    singură dată chiar dacă markerul apare de mai multe ori acolo. Text scris
    ÎNAINTE de primul heading (coperta) primește eticheta convențională de mai
    jos — nu ar trebui să se întâmple în practică (coperta scrie numele
    clientului, nu texte libere), dar tot trebuie să aibă o etichetă, nu una
    lipsă."""
    capitol_curent = "(înainte de primul capitol)"
    capitole_vazute: set[str] = set()
    capitole: list[str] = []
    numar_total = 0

    def _numara(text: str) -> None:
        nonlocal numar_total
        aparitii = text.count(scope.PLACEHOLDER)
        if not aparitii:
            return
        numar_total += aparitii
        if capitol_curent not in capitole_vazute:
            capitole_vazute.add(capitol_curent)
            capitole.append(capitol_curent)

    for bloc in _blocuri_document(doc):
        if isinstance(bloc, Paragraph):
            stil_nume = bloc.style.name if bloc.style is not None else ""
            if stil_nume in ("Heading 1", "Heading 2") and bloc.text.strip():
                capitol_curent = bloc.text.strip()
            _numara(bloc.text)
        elif isinstance(bloc, Table):
            for rand in bloc.rows:
                # O celulă îmbinată orizontal apare de mai multe ori în
                # `rand.cells` (câte o dată per coloană de grilă acoperită) —
                # fără deduplicare după celula XML reală, un singur marker
                # s-ar număra de mai multe ori.
                vazute_tc: set[int] = set()
                for celula in rand.cells:
                    if id(celula._tc) in vazute_tc:
                        continue
                    vazute_tc.add(id(celula._tc))
                    _numara(celula.text)

    return numar_total, capitole


async def run_scope_document_pipeline(
    gazda_path: Path | None,
    cfg: dict,
    continut: dict,
    on_step=None,
) -> tuple[Path, dict]:
    """Construiește documentul complet de scop (11 capitole) — `skills.scop_core.scope`.

    `gazda_path` e opțională: fără ea, `scope.genereaza` (prin
    `stil.document_din_gazda(None)`) degradează la un document nou, cu
    stilurile implicite din python-docx — nu ridică eroare de fișier lipsă.

    `continut` — dicționarul brut din browser (`client`, `suplimentare`,
    `acoperire`, `delimitare`, `flux_operational`, `ordine_cap4`, `beneficii`,
    `confirmari`) — se convertește AICI, defensiv, în dataclass-urile din
    `scope.py`; vezi funcțiile `_*_din_dict(uri)` de mai sus pentru regulile
    de toleranță pe fiecare bucată. `cfg` — configurarea structurală
    (document, document_frate, sectiuni_core, capitole, stil) — trece prin
    `_cfg_pregatit`, aceeași filosofie defensivă.

    Cele două erori deliberate ale lui `scope.genereaza` (element suplimentar
    atașat unui modul exclus; ordinea capitolului 4 incompletă sau cu element
    necunoscut) NU sunt prinse aici — traversează neschimbate până la
    apelant, ca routerul să le transforme într-un 422 cu mesaj clar (mesajele
    lor nu conțin nicio cale de fișier — sigure de arătat direct).

    Întoarce `(document_path, sumar)`. `sumar` alimentează UI-ul: capitolele
    scrise, modulele CORE incluse, elementele suplimentare plasate, cerințele
    pe fiecare încadrare A/P/D/N, avertismentele (rânduri respinse la
    conversie, nu erorile deliberate de mai sus) și `placeholder_numar`/
    `placeholder_capitole` — rezultatul lui `_gaseste_placeholdere` pe
    documentul chiar scris, plasa de siguranță din docstring-ul de sus al
    fișierului; un `placeholder_numar` nenul are mereu și un mesaj dedicat,
    primul, în `avertisment`."""
    if on_step:
        on_step("parsing")

    continut = continut if isinstance(continut, dict) else {}

    client = _client_din_dict(continut.get("client"))
    suplimentare, suplimentare_respinse = _suplimentare_din_dicturi(continut.get("suplimentare"))
    acoperire, cerinte_respinse, cerinte_corectate = _acoperire_din_dicturi(continut.get("acoperire"))
    delimitare, delimitare_respinse = _randuri_3col_din_dicturi(
        continut.get("delimitare"), ("zona", "tratat_in", "interfatare"))
    flux_operational, flux_respinse = _randuri_3col_din_dicturi(
        continut.get("flux_operational"), ("etapa", "ce_se_intampla", "rezultat"))
    ordine_cap4 = _ordine_cap4_din_lista(continut.get("ordine_cap4"))
    beneficii = _lista_stringuri(continut.get("beneficii"))
    confirmari, confirmari_respinse = _confirmari_din_dicturi(continut.get("confirmari"))

    cfg_pregatit = _cfg_pregatit(cfg, gazda_path)

    if on_step:
        on_step("building")

    document_path = _mktemp_path(".docx")
    try:
        scope.genereaza(
            cfg_pregatit, document_path, client=client, suplimentare=suplimentare,
            beneficii=beneficii, confirmari=confirmari, acoperire=acoperire,
            delimitare=delimitare, flux_operational=flux_operational, ordine_cap4=ordine_cap4,
        )
    except Exception:
        # Include cele două erori deliberate (ValueError) — se curăță fișierul
        # temporar oricum, dar excepția urcă NESCHIMBATĂ (vezi docstring-ul).
        document_path.unlink(missing_ok=True)
        raise

    doc_citit = docx.Document(str(document_path))
    capitole = [
        p.text.strip() for p in doc_citit.paragraphs
        if p.style is not None and p.style.name == "Heading 1" and p.text.strip()
    ]
    placeholder_numar, placeholder_capitole = _gaseste_placeholdere(doc_citit)

    sc = cfg_pregatit.get("sectiuni_core", {})
    sectiuni = capitol.alege_sectiuni(
        fara=None if sc.get("toate", True) else (sc.get("fara") or None),
        doar=None if sc.get("toate", True) else (sc.get("doar") or None),
    )

    cerinte_pe_incadrare = {"A": 0, "P": 0, "D": 0, "N": 0}
    for c in acoperire:
        cerinte_pe_incadrare[c.incadrare] += 1

    suplimentare_brute = continut.get("suplimentare")
    acoperire_brute = continut.get("acoperire")

    avertismente: list[str] = []
    if placeholder_numar:
        avertismente.append(
            f"ATENȚIE — documentul conține {placeholder_numar} marcaj(e) „{scope.PLACEHOLDER}” "
            "necompletat(e); nu îl trimite clientului așa. Apare în: "
            + "; ".join(placeholder_capitole) + "."
        )
    if suplimentare_respinse:
        avertismente.append(
            f"{suplimentare_respinse} element(e) suplimentar(e) primite din browser nu au putut fi "
            "folosite (titlu lipsă) și au fost ignorate."
        )
    if cerinte_respinse:
        avertismente.append(
            f"{cerinte_respinse} cerință(e) de acoperire primite din browser nu au putut fi folosite "
            "(cerință sau răspuns lipsă) și au fost ignorate."
        )
    if cerinte_corectate:
        avertismente.append(
            f"{cerinte_corectate} cerință(e) de acoperire aveau o încadrare necunoscută și au fost "
            "trecute la «D. De definit», ca să nu dispară tăcut din analiză."
        )
    if delimitare_respinse:
        avertismente.append(
            f"{delimitare_respinse} rând(uri) din tabelul de delimitare față de documentul-frate nu "
            "au putut fi folosite (zonă, tratare sau interfațare lipsă) și au fost ignorate."
        )
    if flux_respinse:
        avertismente.append(
            f"{flux_respinse} etapă(e) din fluxul operațional nu au putut fi folosite (etapă, "
            "descriere sau rezultat lipsă) și au fost ignorate."
        )
    if confirmari_respinse:
        avertismente.append(
            f"{confirmari_respinse} punct(e) de confirmat nu au putut fi folosite (aspect sau motiv "
            "lipsă) și au fost ignorate."
        )

    sumar = {
        "capitole_scrise": len(capitole),
        "capitole": capitole,
        "module_core": len(sectiuni),
        "elemente_suplimentare_primite": len(suplimentare_brute) if isinstance(suplimentare_brute, list) else 0,
        "elemente_suplimentare_plasate": len(suplimentare),
        "elemente_suplimentare_respinse": suplimentare_respinse,
        "cerinte_primite": len(acoperire_brute) if isinstance(acoperire_brute, list) else 0,
        "cerinte_plasate": len(acoperire),
        "cerinte_respinse": cerinte_respinse,
        "cerinte_pe_incadrare": cerinte_pe_incadrare,
        "delimitare_randuri_plasate": len(delimitare),
        "delimitare_randuri_respinse": delimitare_respinse,
        "flux_operational_randuri_plasate": len(flux_operational),
        "flux_operational_randuri_respinse": flux_respinse,
        "confirmari_plasate": len(confirmari),
        "confirmari_respinse": confirmari_respinse,
        "placeholder_numar": placeholder_numar,
        "placeholder_capitole": placeholder_capitole,
        "avertisment": " ".join(avertismente),
    }
    return document_path, sumar

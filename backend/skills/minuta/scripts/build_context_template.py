# -*- coding: utf-8 -*-
"""Genereaza template-ul de context pentru minuta (.docx).

Contextul e informatia STABILA a proiectului — cea pe care transcriptul nu o
contine, dar de care depinde calitatea minutei: cine sunt oamenii si ce rol au,
ce inseamna abrevierile pe care transcrierea automata le stalceste, ce s-a decis
inainte. Fisierul se completeaza o data per proiect si se reutilizeaza la fiecare
sedinta.

Fiecare capitol are un titlu (Heading) + un tabel sau paragrafe completabile,
astfel incat parserul sa poata citi ce a completat utilizatorul.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x70, 0x70, 0x70)


def _hint(doc: Document, text: str) -> None:
    """Instructiune pentru cel care completeaza — parserul o ignora."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def _table(doc: Document, headers: list[str], rows: int, widths: list[int] | None = None):
    t = doc.add_table(rows=1 + rows, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
    return t


def build_context_template(output_path: Path) -> Path:
    doc = Document()

    title = doc.add_heading("Context Proiect — pentru generarea minutelor", level=0)
    for run in title.runs:
        run.font.color.rgb = NAVY

    intro = doc.add_paragraph()
    r = intro.add_run(
        "Completați o singură dată per proiect și reutilizați la fiecare ședință. "
        "Toate capitolele sunt opționale — completați doar ce vă e util; ce lăsați gol "
        "este ignorat. Nu ștergeți titlurile capitolelor."
    )
    r.font.size = Pt(10)
    r.italic = True

    # 1 — Identificare
    doc.add_heading("1. Proiect și client", level=1)
    _hint(doc, "Apare în antetul minutei și ajută la denumirea corectă a părților.")
    t = _table(doc, ["Câmp", "Valoare"], 5)
    for i, label in enumerate(
        ["Client", "Proiect", "Cod proiect", "Faza / etapa", "Manager de proiect"], start=1
    ):
        t.rows[i].cells[0].text = label

    # 2 — Participanti
    doc.add_heading("2. Participanți recurenți", level=1)
    _hint(
        doc,
        "Cea mai utilă secțiune: transcrierea dă doar nume, nu și rolul. Cu tabelul de mai jos, "
        "minuta atribuie corect responsabilitățile și scrie numele complet chiar dacă în "
        "transcriere apare prescurtat.",
    )
    t = _table(doc, ["Nume", "Rol", "Organizație"], 8)

    # 3 — Glosar
    doc.add_heading("3. Glosar și terminologie", level=1)
    _hint(
        doc,
        "Termenii pe care transcrierea automată îi scrie greșit (module, abrevieri, denumiri de "
        "ecrane). Cu ei, minuta folosește forma corectă în loc să reproducă eroarea.",
    )
    # Fara exemple pre-completate in tabel: un template necompletat trebuie sa
    # produca un context GOL, nu unul cu date inventate de noi.
    _hint(doc, "Exemple: NIR = Notă de intrare-recepție · PVD = Proces verbal de diferențe")
    _table(doc, ["Termen / abreviere", "Înseamnă"], 10)

    # 4 — Scop
    doc.add_heading("4. Scopul proiectului și module în lucru", level=1)
    _hint(doc, "Ce se implementează și ce NU este în scop — ajută la separarea subiectelor.")
    doc.add_paragraph("Module / arii în scop:")
    doc.add_paragraph("", style="List Bullet")
    doc.add_paragraph("", style="List Bullet")
    doc.add_paragraph("În afara scopului:")
    doc.add_paragraph("", style="List Bullet")

    # 5 — Istoric decizii
    doc.add_heading("5. Decizii luate anterior", level=1)
    _hint(
        doc,
        "Deciziile deja agreate. Minuta le poate lega de discuția curentă și nu le reraportează "
        "ca fiind noi.",
    )
    _table(doc, ["Data", "Decizie"], 6)

    # 6 — Actiuni deschise
    doc.add_heading("6. Acțiuni deschise din ședințele anterioare", level=1)
    _hint(
        doc,
        "Ce era în lucru înainte. Dacă în ședință se discută despre ele, minuta marchează "
        "explicit ce s-a închis și ce rămâne.",
    )
    _table(doc, ["Responsabil", "Acțiune", "Termen"], 6)

    # 7 — Preferinte
    doc.add_heading("7. Preferințe de redactare", level=1)
    _hint(doc, "Cum vreți să arate minuta. Scrieți liber, în cuvintele voastre.")
    _hint(
        doc,
        "Exemple: Evidențiați riscurile într-o secțiune separată. · "
        "Folosiți „scop”, nu „perimetru”. · "
        "Nu includeți discuțiile despre planificarea internă.",
    )
    doc.add_paragraph("")
    doc.add_paragraph("")

    doc.save(str(output_path))
    return output_path


if __name__ == "__main__":
    out = Path(__file__).parent.parent / "template" / "Context_Proiect_Template.docx"
    build_context_template(out)
    print(f"Salvat: {out}")

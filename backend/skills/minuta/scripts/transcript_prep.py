# -*- coding: utf-8 -*-
"""Pregătește transcrierea Teams (.vtt sau .docx) pentru model: blocuri cu ora și vorbitorul.

Teams exportă câte o replică scurtă pe rând — în VTT cu identificator, interval de timp și
<v Vorbitor>text</v>, în DOCX ca „Nume   0:04” urmat de text. Aici replicile consecutive ale
aceluiași vorbitor se unesc în blocuri „[00:12:30] Nume: text”, iar ezitările pure („ăă”,
„hmm”) se scot. Ora rămâne în text, ca minuta să poată trimite la momentul neclar din
înregistrare ([CLARIFICARE NEEDED: …, 00:12:30]). Pe trei ședințe reale, textul trimis la
Claude a scăzut cu 52–55% față de extragerea anterioară.

Portat din AI_Claude/tools/whisper/vtt_clean.py, fără glosarul de proiect: pe site, termenii
proiectului vin din fișierul de Context Proiect și se corectează în prompt.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document

MERGE_GAP_S = 3.0      # replicile aceluiași vorbitor la mai puțin de 3 s una de alta se unesc…
MERGE_SPAN_S = 60.0    # …dar un bloc nu acoperă mai mult de un minut
LONG_TURN_LINES = 40   # o replică DOCX mai lungă (fără timpi intermediari) se împarte…
SPLIT_LINES = 12       # …în bucăți de câte 12 rânduri

# Doar ezitări pure. „da”, „ok”, „evet”, „tamam” rămân: pot fi răspunsul la o decizie.
HESITATIONS = {
    "ă", "ăă", "ăăă", "î", "îî", "îîî", "ee", "eee", "ı", "ıı", "ııı", "ıııı",
    "hm", "hmm", "hmmm", "mm", "mmm", "mhm", "hıhı", "uh", "um", "uhm", "eh",
}

_TIME = r"(?:(\d+):)?(\d{1,2}):(\d{2})(?:[.,](\d{1,3}))?"
TIMING = re.compile(rf"^{_TIME}\s*-->\s*{_TIME}")
VOICE = re.compile(r"<v(?:\.[^\s>]*)?\s+([^>]+?)\s*>")
TAG = re.compile(r"<[^>]+>")
DOCX_HEAD = re.compile(r"^(.+?)\s{2,}(\d+:\d{2}(?::\d{2})?)\s*$")
# rândurile de sistem Teams din exportul Word: nu sunt vorbire
SYSTEM_LINE = re.compile(r"\b(?:started|stopped) transcription$|\ba (?:pornit|oprit) transcrierea$", re.I)


@dataclass
class Cue:
    start: float
    end: float
    speaker: str
    text: str
    mergeable: bool = True  # False pentru bucățile unei replici DOCX lungi: nu se reunesc


def _secs(h, m, s, ms) -> float:
    return int(h or 0) * 3600 + int(m) * 60 + int(s) + (int(ms.ljust(3, "0")) / 1000 if ms else 0.0)


def hms(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def _speaker(name: str) -> str:
    return re.sub(r"\s+", " ", name).strip()


def read_vtt(path: Path) -> list[Cue]:
    raw = path.read_text(encoding="utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
    cues: list[Cue] = []
    for block in re.split(r"\n\s*\n", raw):
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        for i, line in enumerate(lines):
            m = TIMING.match(line)
            if not m:
                continue
            g = m.groups()
            payload = " ".join(lines[i + 1:])
            voice = VOICE.search(payload)
            text = re.sub(r"\s+", " ", TAG.sub("", payload)).strip()
            if text:
                cues.append(Cue(_secs(*g[:4]), _secs(*g[4:]), _speaker(voice.group(1)) if voice else "", text))
            break
    return cues


def _docx_secs(t: str) -> float:
    parts = [int(x) for x in t.split(":")]
    while len(parts) < 3:
        parts.insert(0, 0)
    return parts[0] * 3600 + parts[1] * 60 + parts[2]


def read_docx(path: Path) -> list[Cue]:
    """Transcrierea Teams exportată în Word. Listă goală dacă documentul nu are formatul Teams."""
    # exportul Teams pune adesea rândurile în același paragraf, separate prin salt de linie
    lines = [ln.strip() for p in Document(str(path)).paragraphs for ln in p.text.split("\n")]
    turns: list[list] = []  # [start, vorbitor, rânduri]
    for ln in lines:
        if not ln or SYSTEM_LINE.search(ln):
            continue
        m = DOCX_HEAD.match(ln)
        if m:
            turns.append([_docx_secs(m.group(2)), _speaker(m.group(1)), []])
        elif turns:
            turns[-1][2].append(ln)
    if not turns:  # o sală înregistrată de pe un singur cont dă o singură replică — e tot format Teams
        return []
    cues: list[Cue] = []
    for i, (start, sp, rows) in enumerate(turns):
        end = max(start, turns[i + 1][0]) if i + 1 < len(turns) else start
        if len(rows) > LONG_TURN_LINES:
            for j in range(0, len(rows), SPLIT_LINES):
                cues.append(Cue(start, end, sp, " ".join(rows[j:j + SPLIT_LINES]), mergeable=False))
        elif rows:
            cues.append(Cue(start, end, sp, " ".join(rows)))
    return cues


def is_hesitation(text: str) -> bool:
    words = re.sub(r"[^\wăâîșțşţçğıöüĂÂÎȘȚŞŢÇĞİÖÜ\s]", " ", text.lower()).split()
    return not words or all(w in HESITATIONS for w in words)


def merge(cues: list[Cue]) -> list[Cue]:
    out: list[Cue] = []
    for c in cues:
        if is_hesitation(c.text):
            continue
        last = out[-1] if out else None
        if (last and last.speaker == c.speaker and last.mergeable and c.mergeable
                and c.start - last.end < MERGE_GAP_S and c.end - last.start < MERGE_SPAN_S):
            # Teams repetă uneori aceeași replică; un „da” scurt care coincide nu e o repetare
            if not (len(c.text.split()) >= 3 and last.text.endswith(c.text)):
                last.text += " " + c.text
            last.end = max(last.end, c.end)
        else:
            out.append(Cue(c.start, c.end, c.speaker, c.text, c.mergeable))
    return out


def _plain_text(path: Path) -> str:
    """Rezerva pentru fișierele care nu au formatul Teams: textul ca atare."""
    if path.suffix.lower() == ".vtt":
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        return "\n".join(
            TAG.sub("", ln).strip() for ln in lines
            if ln.strip() and ln.strip() != "WEBVTT" and "-->" not in ln
        )
    return "\n".join(p.text for p in Document(str(path)).paragraphs if p.text.strip())


# ora de la începutul unui bloc, „[00:12:30] ”
_ORA_BLOC = re.compile(r"^\[(\d{2}:\d{2}:\d{2})\]", re.M)


def prepare_transcripts(fisiere: list[tuple[Path, str]]) -> str:
    """Una sau mai multe înregistrări ale aceleiași zile, într-o singură transcriere.

    Teams împarte adesea o zi de ședințe în mai multe înregistrări (o pauză, altă sală, altă
    ședință cu aceeași companie). Fiecare înregistrare primește un antet cu numele ei, iar ora din
    blocuri primește în față numărul înregistrării — „[(2) 00:12:03]” —, ca minuta să poată trimite
    la momentul exact chiar dacă ora se repetă de la o înregistrare la alta. O singură înregistrare
    rămâne exact ca la `prepare_transcript`.
    """
    if len(fisiere) == 1:
        return prepare_transcript(fisiere[0][0])
    parti = []
    for k, (cale, nume) in enumerate(fisiere, 1):
        text = _ORA_BLOC.sub(lambda m: f"[({k}) {m.group(1)}]", prepare_transcript(cale))
        parti.append(f"=== Înregistrarea {k} din {len(fisiere)}: „{nume}” ===\n{text}")
    return "\n\n".join(parti)


def prepare_transcript(path: Path) -> str:
    """Textul de dat modelului: un bloc pe rând, „[hh:mm:ss] Vorbitor: text”."""
    cues = read_vtt(path) if path.suffix.lower() == ".vtt" else read_docx(path)
    if not cues:
        return _plain_text(path)
    return "\n".join(
        f"[{hms(c.start)}] {c.speaker}: {c.text}" if c.speaker else f"[{hms(c.start)}] {c.text}"
        for c in merge(cues)
    )

# -*- coding: utf-8 -*-
"""Genereaza livrabilele: agenda (Word) si planul cu participanti (Excel)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F3864"
BLUE_MID = "2E5496"
BLUE_LIGHT = "F2F6FB"
YELLOW_FILL = "FFF9E6"
GREY_BORDER = "BFBFBF"
WHITE = "FFFFFF"
GREY_TEXT = "707070"

ZILE_NUME = ["Ziua 1", "Ziua 2", "Ziua 3", "Ziua 4", "Ziua 5",
             "Ziua 6", "Ziua 7", "Ziua 8", "Ziua 9", "Ziua 10"]


def fmt_ore(ore: float) -> str:
    return f"{ore:g}".replace(".", ",") + "h"


# ─────────────────────────────────────────────────────────── Word

def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _borders(cell) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), GREY_BORDER)
        b.append(el)
    tc_pr.append(b)


def _heading(doc, text: str, size: int, color: str, level: int | None = None):
    p = doc.add_paragraph()
    if level:
        try:
            p.style = doc.styles[f"Heading {level}"]
        except KeyError:
            pass
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p, r


def build_word(plan: dict, meta: dict, path: Path) -> Path:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Agenda Training")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["titlu_program"])
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(BLUE_MID)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["client"])
    r.bold = True
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{len(plan['zile'])} zile · {fmt_ore(plan['total_ore'])} efectiv")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GREY_TEXT)

    # Tabel de efort — raspunde direct la „cat alocam fiecarui modul"
    doc.add_paragraph()
    _heading(doc, "Efort estimat pe module", 13, NAVY, level=1)

    rows = [(z, m) for z in plan["zile"] for m in z["module"]]
    t = doc.add_table(rows=1 + len(rows) + 1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Nr.", "Modul", "Zi", "Efort"]):
        c = t.rows[0].cells[i]
        c.text = h
        _shade(c, NAVY)
        _borders(c)
        run = c.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(10)

    for i, (z, m) in enumerate(rows, start=1):
        for col, val in enumerate([
            str(m["nr"]), m["nume"], ZILE_NUME[z["zi"] - 1], fmt_ore(m["ore"])
        ]):
            c = t.rows[i].cells[col]
            c.text = val
            _borders(c)
            if i % 2 == 0:
                _shade(c, BLUE_LIGHT)
            c.paragraphs[0].runs[0].font.size = Pt(10)

    last = t.rows[len(rows) + 1]
    for col, val in enumerate(["", "TOTAL", f"{len(plan['zile'])} zile",
                               fmt_ore(plan["total_ore"])]):
        c = last.cells[col]
        c.text = val
        _borders(c)
        _shade(c, BLUE_LIGHT)
        if val:
            c.paragraphs[0].runs[0].bold = True
            c.paragraphs[0].runs[0].font.size = Pt(10)

    for col, w in enumerate([Cm(1.2), Cm(8.5), Cm(2.5), Cm(2.3)]):
        for row in t.rows:
            row.cells[col].width = w

    if plan["excluse"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        r = p.add_run("Neacoperit la această durată: ")
        r.bold = True
        r.font.size = Pt(9.5)
        r2 = p.add_run(", ".join(m["nume"] for m in plan["excluse"]))
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor.from_string("C00000")

    # Agenda pe zile
    for z in plan["zile"]:
        doc.add_paragraph()
        p, _ = _heading(
            doc,
            f"{ZILE_NUME[z['zi'] - 1]}   |   {meta['interval']}   |   {fmt_ore(z['ore'])} efectiv",
            12, NAVY, level=1,
        )
        pf = p.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(2)
        p_pr = p._element.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), NAVY)
        bdr.append(bottom)
        p_pr.insert(0, bdr)

        for m in z["module"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            try:
                p.style = doc.styles["Heading 2"]
            except KeyError:
                pass
            r = p.add_run(f"{m['nr']}.  {m['nume']}")
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor.from_string(BLUE_MID)
            r2 = p.add_run(f"   —   efort: {fmt_ore(m['ore'])}")
            r2.bold = True
            r2.font.size = Pt(10.5)
            r2.font.color.rgb = RGBColor.from_string("000000")

            if m.get("continuare"):
                continue  # subiectele s-au listat in ziua in care a inceput modulul

            for titlu_grup, subiecte in m["grupe"]:
                pg = doc.add_paragraph()
                pg.paragraph_format.left_indent = Cm(0.6)
                pg.paragraph_format.space_before = Pt(4)
                pg.paragraph_format.space_after = Pt(1)
                rg = pg.add_run(titlu_grup)
                rg.bold = True
                rg.font.size = Pt(10)
                if titlu_grup == "Particularități client":
                    rg.font.color.rgb = RGBColor.from_string("B8860B")
                for s in subiecte:
                    ps = doc.add_paragraph(style="List Bullet")
                    ps.paragraph_format.left_indent = Cm(1.3)
                    ps.paragraph_format.space_after = Pt(0)
                    ps.add_run(s).font.size = Pt(10)

    doc.save(str(path))
    return path


# ─────────────────────────────────────────────────────────── Excel

def build_excel(plan: dict, meta: dict, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan training"

    thin = Side(style="thin", color=GREY_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=NAVY)
    head_font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    zi_fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    completat_fill = PatternFill("solid", fgColor=YELLOW_FILL)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = f"Agenda Training — {meta['client']}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A2"] = (f"{meta['titlu_program']} · {len(plan['zile'])} zile · "
                f"{fmt_ore(plan['total_ore'])} efectiv")
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color=GREY_TEXT)

    headers = ["Zi", "Data", "Interval orar", "Nr.", "Modul", "Conținut acoperit",
               "Efort (ore)", "Audiență recomandată", "Participanți client", "Observații"]
    ws.append([])
    ws.append(headers)
    hr = ws.max_row
    for i in range(1, len(headers) + 1):
        c = ws.cell(row=hr, column=i)
        c.fill = head_fill
        c.font = head_font
        c.alignment = center
        c.border = border

    row = hr + 1
    for z in plan["zile"]:
        start = row
        for m in z["module"]:
            continut = " · ".join(t for t, _ in m["grupe"])
            if m.get("nr_particularitati"):
                continut += f"  [+{m['nr_particularitati']} particularități client]"
            vals = [ZILE_NUME[z["zi"] - 1], "<zz.ll.aaaa>", meta["interval"],
                    m["nr"], m["nume"], continut, m["ore"], m["audienta"], "", ""]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.border = border
                c.font = Font(name="Calibri", size=10)
                c.alignment = center if col in (1, 2, 3, 4, 7) else wrap_top
            ws.cell(row=row, column=7).number_format = "0.0"
            ws.cell(row=row, column=9).fill = completat_fill
            row += 1
        if row - start > 1:
            for col in (1, 2, 3):
                ws.merge_cells(start_row=start, start_column=col,
                               end_row=row - 1, end_column=col)
                ws.cell(row=start, column=col).alignment = center
        for col in (1, 2, 3):
            ws.cell(row=start, column=col).fill = zi_fill

    total_row = row
    ws.cell(row=total_row, column=5, value="TOTAL").font = Font(
        name="Calibri", size=10, bold=True)
    tc = ws.cell(row=total_row, column=7, value=f"=SUM(G{hr + 1}:G{total_row - 1})")
    tc.font = Font(name="Calibri", size=10, bold=True)
    tc.number_format = "0.0"
    tc.alignment = center
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).border = border
        ws.cell(row=total_row, column=col).fill = zi_fill

    for i, w in enumerate([9, 12, 14, 5, 32, 50, 11, 26, 30, 24], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(hr + 1, total_row):
        ws.row_dimensions[r].height = 46
    ws.freeze_panes = f"A{hr + 1}"
    ws.auto_filter.ref = f"A{hr}:J{total_row - 1}"

    # Foaia 2 — cine participa la ce
    ws2 = wb.create_sheet("Participanți")
    ws2["A1"] = "Participanți din partea clientului"
    ws2["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws2["A2"] = ("Se completează de client. Bifați modulele la care participă fiecare "
                 "persoană — nu toți utilizatorii au nevoie de toate modulele.")
    ws2["A2"].font = Font(name="Calibri", size=10, italic=True, color=GREY_TEXT)

    module_unice = []
    vazute = set()
    for z in plan["zile"]:
        for m in z["module"]:
            if m["nr"] not in vazute:
                vazute.add(m["nr"])
                module_unice.append(m)

    head2 = ["Nr.", "Nume și prenume", "Funcție", "Departament", "Email"] + [
        f"M{m['nr']}. {m['nume'].replace('Modul ', '')[:26]}" for m in module_unice]
    ws2.append([])
    ws2.append(head2)
    h2 = ws2.max_row
    for i in range(1, len(head2) + 1):
        c = ws2.cell(row=h2, column=i)
        c.fill = head_fill
        c.font = head_font
        c.alignment = center
        c.border = border

    dv = DataValidation(type="list", formula1='"DA,NU"', allow_blank=True)
    ws2.add_data_validation(dv)
    for i in range(1, 21):
        r = h2 + i
        ws2.cell(row=r, column=1, value=i).alignment = center
        for col in range(1, len(head2) + 1):
            c = ws2.cell(row=r, column=col)
            c.border = border
            c.font = Font(name="Calibri", size=10)
            if col > 5:
                c.alignment = center
                c.fill = completat_fill
        dv.add(f"F{r}:{get_column_letter(len(head2))}{r}")

    for i, w in enumerate([5, 26, 22, 20, 26] + [14] * len(module_unice), start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.row_dimensions[h2].height = 44
    ws2.freeze_panes = f"F{h2 + 1}"

    wb.save(str(path))
    return path
